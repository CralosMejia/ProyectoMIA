"""Lector de la plantilla oficial de carga (CU-01, D4-BE-01) con ``openpyxl``.

Abre el libro en memoria con ``read_only=True`` y ``data_only=True``, comprueba la hoja
``Registros`` y sus titulos (fila 1, columnas A..N) y entrega los valores **crudos** de
A..N de cada fila con algun dato (fechas como ``date``/``datetime``, numeros tal cual); la
validacion por fila es del dominio. Las columnas O y P (formulas de apoyo) no se leen ni se
validan. Las filas vacias intermedias se saltan conservando el numero de fila real.
Los ``UserWarning`` de ``openpyxl`` (p. ej. "Cell G2 ... outside the limits for dates") citan
coordenada y valor de celda: se silencian solo durante la lectura (T6-OBS-1).

Sin hoja ``Registros`` y con alguna hoja cuyo nombre es un mes -> formato antiguo de hojas
mensuales. Los errores nunca incluyen valores de celdas y el lector no escribe en disco ni
en logs.
"""

import io
import unicodedata
import warnings
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from typing import Protocol, cast

from openpyxl import load_workbook  # type: ignore[import-untyped]  # openpyxl no publica stubs
from openpyxl.utils.exceptions import (  # type: ignore[import-untyped]  # sin stubs
    InvalidFileException,
)
from openpyxl.workbook.workbook import (  # type: ignore[import-untyped]  # sin stubs
    Workbook,
)

from app.domain.cargas.errores import MotivoPlantillaInvalida, PlantillaInvalidaError
from app.domain.cargas.fila import COLUMNAS, PRIMERA_FILA_DATOS, FilaPlantilla, PlantillaLeida
from app.infrastructure.historico.alias_columnas import MESES

MAX_FILAS_PLANTILLA = 500
# S3 (R-013): limites contra libros "bomba" (zip muy comprimido o hojas con miles de filas
# vacias). Se comprueban con ``zipfile`` antes de abrir el libro con ``openpyxl``.
MAX_BYTES_DESCOMPRIMIDOS = 50 * 1024 * 1024
MAX_RELACION_COMPRESION = 100
MAX_FILAS_ITERADAS = 5000
HOJA_REGISTROS = "Registros"
TITULOS_PLANTILLA: tuple[str, ...] = (
    "ESTADO DEL REGISTRO",
    "NOMBRES Y APELLIDOS",
    "CÉDULA",
    "CELULAR",
    "CORREO",
    "TIPO DE PLAN",
    "FECHA DE INICIO",
    "FECHA DE VENCIMIENTO",
    "FORMA DE PAGO",
    "N.º DE TRANSACCIÓN",
    "MONTO A FACTURAR (USD)",
    "MONTO COBRADO (USD)",
    "ASISTENCIAS EN EL PERÍODO",
    "OBSERVACIONES",
)
"""Titulos oficiales de A..N (mayusculas y tildes exactas; se comparan tras ``strip()``)."""

_NUMERO_COLUMNAS = len(COLUMNAS)
_ERRORES_DE_ARCHIVO = (
    zipfile.BadZipFile,
    InvalidFileException,
    KeyError,
    OSError,
    ValueError,
    TypeError,
    EOFError,
)


class _Hoja(Protocol):
    """Lo unico que el lector usa de una hoja de ``openpyxl``."""

    def iter_rows(
        self, *, min_row: int, max_col: int, values_only: bool
    ) -> Iterator[tuple[object, ...]]: ...


def _texto_titulo(valor: object) -> str:
    texto = "" if valor is None else str(valor)
    return unicodedata.normalize("NFC", texto.strip())


def _vacia(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _es_hoja_mensual(nombre: str) -> bool:
    return nombre.strip().upper() in MESES


def _comprobar_tamano_descomprimido(contenido: bytes) -> None:
    """Rechaza el zip si declara > 50 MB descomprimidos o una relacion de compresion > 100."""
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as archivo:
            infos = archivo.infolist()
    except _ERRORES_DE_ARCHIVO as error:
        raise PlantillaInvalidaError(MotivoPlantillaInvalida.DANADO) from error
    descomprimido = sum(info.file_size for info in infos)
    comprimido = max(sum(info.compress_size for info in infos), 1)
    if (
        descomprimido > MAX_BYTES_DESCOMPRIMIDOS
        or descomprimido / comprimido > MAX_RELACION_COMPRESION
    ):
        raise PlantillaInvalidaError(MotivoPlantillaInvalida.DEMASIADO_GRANDE)


class LectorPlantillaOpenpyxl:
    """Implementa ``LectorPlantilla``."""

    def leer(self, contenido: bytes) -> PlantillaLeida:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
            return self._leer(contenido)

    def _leer(self, contenido: bytes) -> PlantillaLeida:
        libro = self._abrir(contenido)
        try:
            hoja = self._hoja_registros(libro)
            filas = hoja.iter_rows(min_row=1, max_col=_NUMERO_COLUMNAS, values_only=True)
            try:
                return PlantillaLeida(self._leer_filas(filas))
            except PlantillaInvalidaError:
                raise
            except _ERRORES_DE_ARCHIVO as error:
                raise PlantillaInvalidaError(MotivoPlantillaInvalida.DANADO) from error
        finally:
            libro.close()

    @staticmethod
    def _abrir(contenido: bytes) -> Workbook:
        _comprobar_tamano_descomprimido(contenido)
        try:
            libro: Workbook = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except _ERRORES_DE_ARCHIVO as error:
            raise PlantillaInvalidaError(MotivoPlantillaInvalida.DANADO) from error
        return libro

    @staticmethod
    def _hoja_registros(libro: Workbook) -> _Hoja:
        nombres: list[str] = list(libro.sheetnames)
        for nombre in nombres:
            if nombre.strip().casefold() == HOJA_REGISTROS.casefold():
                return cast(_Hoja, libro[nombre])
        if any(_es_hoja_mensual(nombre) for nombre in nombres):
            raise PlantillaInvalidaError(MotivoPlantillaInvalida.FORMATO_ANTIGUO)
        raise PlantillaInvalidaError(MotivoPlantillaInvalida.SIN_REGISTROS)

    @staticmethod
    def _leer_filas(filas: Iterable[Sequence[object]]) -> tuple[FilaPlantilla, ...]:
        iterador = iter(filas)
        titulos = next(iterador, ())
        leidos = tuple(_texto_titulo(v) for v in titulos)
        leidos += ("",) * (_NUMERO_COLUMNAS - len(leidos))
        if leidos[:_NUMERO_COLUMNAS] != TITULOS_PLANTILLA:
            raise PlantillaInvalidaError(MotivoPlantillaInvalida.TITULOS)

        resultado: list[FilaPlantilla] = []
        for numero, valores in enumerate(iterador, start=PRIMERA_FILA_DATOS):
            if numero - PRIMERA_FILA_DATOS >= MAX_FILAS_ITERADAS:
                raise PlantillaInvalidaError(MotivoPlantillaInvalida.DEMASIADO_GRANDE)
            crudos = tuple(valores[:_NUMERO_COLUMNAS])
            if all(_vacia(v) for v in crudos):
                continue
            if len(resultado) == MAX_FILAS_PLANTILLA:
                raise PlantillaInvalidaError(MotivoPlantillaInvalida.DEMASIADAS_FILAS)
            resultado.append(FilaPlantilla.desde_valores(numero, crudos))
        return tuple(resultado)
