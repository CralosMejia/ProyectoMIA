"""Excel de la lista priorizada con openpyxl, generado en memoria (CU-12).

- Igual que la pantalla (R-008, decisiones (a)-(c)): hojas ``Socios prioritarios a
  contactar`` (items ``dentro_top_k``), ``Otros socios por vencer`` (el resto; si no hay,
  solo encabezados) y ``Resumen`` con terminos amigables y el nombre de la lista. Las dos
  primeras por ``prioridad`` ascendente y con las columnas de la tabla de la pantalla: sin
  celular, score ni monto.
- R-009 (j): columnas ``Gestión`` (``Sin gestión`` / ``Pendiente`` / ``Cerrada``) y
  ``Campaña`` (nombre guardado en la gestion; ``Sin campaña`` si la gestion no tiene; ``—`` si
  no hay gestion), con el estado al momento de descargar. El boton de accion no se exporta.
- Motivos sin el aporte numerico final entre parentesis (misma regla que ``APORTE_FINAL`` de
  ``TablaLista.tsx``); sin motivos -> ``SIN_FACTORES``.
- RN-12.1: todo ocurre en ``BytesIO``. openpyxl escribe cada hoja en un archivo temporal;
  aqui se sustituye ese paso por un buffer en memoria, asi que no se toca el disco.
- RN-12.2: un texto que empieza por ``=``, ``+``, ``-`` o ``@`` (incluido el nombre de la
  lista) se guarda como texto con el prefijo ``'``; ninguna celda es una formula.
- La cedula se guarda siempre como texto (conserva el cero inicial); el correo no se exporta.
- RN-12.3: solo usa la lista guardada (``ListaDetalle``); no ejecuta el modelo.
- Determinista: las propiedades del libro y las fechas de las entradas del zip salen de
  ``creada_en``, asi que la misma lista produce siempre los mismos bytes.
"""

import io
import re
import zipfile
from collections import Counter
from collections.abc import Sequence
from datetime import UTC, datetime, timedelta
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]  # openpyxl no publica stubs
from openpyxl.cell.cell import (  # type: ignore[import-untyped]  # openpyxl no publica stubs
    ILLEGAL_CHARACTERS_RE,
)
from openpyxl.drawing.spreadsheet_drawing import (  # type: ignore[import-untyped]  # sin stubs
    SpreadsheetDrawing,
)
from openpyxl.worksheet._writer import (  # type: ignore[import-untyped]  # sin stubs
    WorksheetWriter,
)
from openpyxl.writer.excel import ExcelWriter  # type: ignore[import-untyped]  # sin stubs

from app.application.listas.ports import ItemLista, ListaDetalle
from app.domain.common.enums import EstadoGestion, NivelRiesgo

COLUMNAS: tuple[str, ...] = (
    "Prioridad",
    "Nombre",
    "Cédula",
    "Plan",
    "Vence el",
    "Días para vencer",
    "Nivel de riesgo",
    "Motivos principales",
    "Gestión",
    "Campaña",
)
SIN_GESTION = "Sin gestión"
SIN_CAMPANIA = "Sin campaña"
SIN_DATO = "—"
GESTION_TEXTO: dict[EstadoGestion, str] = {
    EstadoGestion.PENDIENTE: "Pendiente",
    EstadoGestion.CERRADA: "Cerrada",
}
SIN_FACTORES = "Sin factores de riesgo destacados"
HOJA_PRIORITARIOS = "Socios prioritarios a contactar"
HOJA_OTROS = "Otros socios por vencer"
HOJA_RESUMEN = "Resumen"
FORMATO_FECHA = "dd/mm/yyyy"
NIVEL_TEXTO: dict[NivelRiesgo, str] = {
    NivelRiesgo.ALTO: "Alto",
    NivelRiesgo.MEDIO: "Medio",
    NivelRiesgo.BAJO: "Bajo",
}
# Aporte numerico al final del texto de un motivo, p. ej. " (+0,12)" (APORTE_FINAL del front).
APORTE_FINAL = re.compile(r"\s*\([^()]*\)\s*$")
_PREFIJOS_FORMULA = ("=", "+", "-", "@")
_SEPARADOR_FACTORES = "; "


def neutralizar(texto: str) -> str:
    """Texto seguro para Excel: quita caracteres de control y desactiva formulas (RN-12.2)."""
    limpio: str = ILLEGAL_CHARACTERS_RE.sub("", texto)
    return "'" + limpio if limpio.startswith(_PREFIJOS_FORMULA) else limpio


def sin_aporte(factor: str) -> str:
    """``"Plan: TRIMESTRAL (+0,09)"`` -> ``"Plan: TRIMESTRAL"``."""
    return APORTE_FINAL.sub("", factor, count=1)


def motivos(factores: Sequence[str]) -> str:
    """Motivos unidos con ``"; "`` y sin aporte; sin motivos -> ``SIN_FACTORES``."""
    return _SEPARADOR_FACTORES.join(sin_aporte(f) for f in factores) or SIN_FACTORES


def gestion_y_campania(item: ItemLista) -> tuple[str, str]:
    """Textos de las columnas ``Gestión`` y ``Campaña``, los mismos que la pantalla."""
    if item.gestion_estado is None:
        return SIN_GESTION, SIN_DATO
    return GESTION_TEXTO[item.gestion_estado], item.campania_asignada or SIN_CAMPANIA


def _utc_sin_zona(instante: datetime) -> datetime:
    return instante.astimezone(UTC).replace(tzinfo=None)


class _ZipDeterminista(zipfile.ZipFile):
    """Zip cuyas entradas llevan siempre la misma fecha (no la hora del sistema)."""

    def __init__(self, destino: io.BytesIO, fecha: datetime) -> None:
        super().__init__(destino, "w", zipfile.ZIP_DEFLATED, allowZip64=True)
        self._fecha = (fecha.year, fecha.month, fecha.day, fecha.hour, fecha.minute, fecha.second)

    def writestr(
        self,
        zinfo_or_arcname: str | zipfile.ZipInfo,
        data: Any,  # noqa: ANN401 - misma firma que ZipFile.writestr (bytes o str)
        compress_type: int | None = None,
        compresslevel: int | None = None,
    ) -> None:
        if isinstance(zinfo_or_arcname, zipfile.ZipInfo):
            info = zinfo_or_arcname
        else:
            info = zipfile.ZipInfo(zinfo_or_arcname, date_time=self._fecha)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
        super().writestr(info, data, compress_type, compresslevel)


class _EscritorEnMemoria(ExcelWriter):  # type: ignore[misc]  # base sin tipos (openpyxl)
    """``ExcelWriter`` que serializa cada hoja en memoria en lugar de un temporal."""

    def write_worksheet(self, ws: Any) -> None:  # noqa: ANN401 - Worksheet de openpyxl sin tipos
        ws._drawing = SpreadsheetDrawing()
        ws._drawing.charts = ws._charts
        ws._drawing.images = ws._images
        writer = WorksheetWriter(ws, out=io.BytesIO())
        writer.write()
        ws._rels = writer._rels
        self._archive.writestr(ws.path[1:], writer.read())
        self.manifest.append(ws)


class ExcelListaOpenpyxl:
    """Implementa ``ExportadorLista``."""

    def generar(self, lista: ListaDetalle) -> bytes:
        creada_en = _utc_sin_zona(lista.creada_en)
        libro = Workbook()
        libro.properties.created = creada_en
        libro.properties.modified = creada_en

        ordenados = sorted(lista.items, key=lambda i: i.prioridad)
        prioritarios = [i for i in ordenados if i.dentro_top_k]
        otros = [i for i in ordenados if not i.dentro_top_k]
        hoja_prioritarios = libro.active
        hoja_prioritarios.title = HOJA_PRIORITARIOS
        self._escribir_items(hoja_prioritarios, lista, prioritarios)
        self._escribir_items(libro.create_sheet(HOJA_OTROS), lista, otros)
        self._escribir_resumen(libro.create_sheet(HOJA_RESUMEN), lista, prioritarios, otros)

        destino = io.BytesIO()
        with _ZipDeterminista(destino, creada_en) as archivo:
            _EscritorEnMemoria(libro, archivo).write_data()
        return destino.getvalue()

    @staticmethod
    def _texto(hoja: Any, fila: int, columna: int, valor: str | None) -> None:  # noqa: ANN401 - Worksheet sin tipos
        celda = hoja.cell(row=fila, column=columna)
        if valor is not None:
            celda.value = neutralizar(valor)
            celda.data_type = "s"

    @staticmethod
    def _fecha(hoja: Any, fila: int, columna: int, valor: object) -> None:  # noqa: ANN401 - Worksheet sin tipos
        celda = hoja.cell(row=fila, column=columna, value=valor)
        celda.number_format = FORMATO_FECHA

    def _escribir_items(self, hoja: Any, lista: ListaDetalle, items: Sequence[ItemLista]) -> None:  # noqa: ANN401 - Worksheet sin tipos
        for columna, encabezado in enumerate(COLUMNAS, start=1):
            self._texto(hoja, 1, columna, encabezado)
        for fila, item in enumerate(items, start=2):
            hoja.cell(row=fila, column=1, value=item.prioridad)
            self._texto(hoja, fila, 2, item.socio_nombre)
            self._texto(hoja, fila, 3, item.cedula)
            self._texto(hoja, fila, 4, item.plan.value)
            self._fecha(hoja, fila, 5, item.fecha_vencimiento)
            hoja.cell(
                row=fila, column=6, value=(item.fecha_vencimiento - lista.fecha_generacion).days
            )
            self._texto(hoja, fila, 7, NIVEL_TEXTO[item.nivel_riesgo])
            self._texto(hoja, fila, 8, motivos(item.factores_principales))
            gestion, campania = gestion_y_campania(item)
            self._texto(hoja, fila, 9, gestion)
            self._texto(hoja, fila, 10, campania)

    def _escribir_resumen(
        self,
        hoja: Any,  # noqa: ANN401 - Worksheet sin tipos
        lista: ListaDetalle,
        prioritarios: Sequence[ItemLista],
        otros: Sequence[ItemLista],
    ) -> None:
        niveles = Counter(item.nivel_riesgo for item in lista.items)
        self._texto(hoja, 1, 1, "Nombre de la lista")
        self._texto(hoja, 1, 2, lista.nombre)
        self._texto(hoja, 2, 1, "Fecha de la lista")
        self._fecha(hoja, 2, 2, lista.fecha_generacion)
        self._texto(hoja, 3, 1, "Vencen hasta")
        self._fecha(hoja, 3, 2, lista.fecha_generacion + timedelta(days=lista.horizonte_dias))
        conteos: list[tuple[str, int]] = [
            (HOJA_PRIORITARIOS, len(prioritarios)),
            (HOJA_OTROS, len(otros)),
            ("Total de socios por vencer", len(lista.items)),
            ("Riesgo alto", niveles.get(NivelRiesgo.ALTO, 0)),
            ("Riesgo medio", niveles.get(NivelRiesgo.MEDIO, 0)),
            ("Riesgo bajo", niveles.get(NivelRiesgo.BAJO, 0)),
        ]
        for numero, (etiqueta, valor) in enumerate(conteos, start=4):
            self._texto(hoja, numero, 1, etiqueta)
            hoja.cell(row=numero, column=2, value=valor)
