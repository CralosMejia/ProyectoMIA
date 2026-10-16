"""Verificacion local de la lista de septiembre frente al notebook M01 (R-015 T11).

Fuera de la suite. Uso (desde ``Backend/``, con el entorno virtual activo)::

    python scripts/verificar_lista.py <carpeta con los libros Control Gym>
        [--referencia ../salidas] [--fecha 2026-08-31] [--horizonte 30] [--top-k 25]

Compone **en memoria y sin base de datos** el mismo camino que la app: ``LectorLibrosExcel``
+ ``ProcesarHistorico`` -> ``EntrenarModeloInicial`` (``RandomForestPredictor``, conjunto
``CONJUNTO_MODELO``) -> las membresias como las guarda la migracion -> ``GenerarLista`` sobre
una unidad de trabajo en memoria. Lo compara con las salidas del notebook
``Notebookes/MVP.ipynb`` (13.8), que se leen en tiempo de ejecucion de la carpeta git-ignorada
``salidas/`` (nunca se copian a un archivo versionado):

- ``m01/despues/parametros_modelo_operativo.json``: conteos de la cartera, ciclos de
  entrenamiento, variables del modelo, medianas de monto y referencia de los factores;
- ``mvp/lista_operativa_2026-09.xlsx`` (hoja ``cartera_completa``): la lista completa con la
  probabilidad sin redondear, el nivel y los factores de cada socio.

Se comparan **todas** las filas de la lista: probabilidad a 8 decimales (la precision que
guarda la app), nivel (cortes fijos 0,70 / 0,50 en ambos), factores (texto exacto) y orden.
Los socios con la misma probabilidad forman un grupo de empate: el notebook los ordena con un
``sort_values`` no estable (orden arbitrario) y la app por vencimiento y ``membresia_id``, asi
que el orden se compara por grupo (mismo conjunto de socios en las mismas posiciones).

Los codigos de socio (``MTB-`` + 8 hex) son un HMAC-SHA256 con una sal secreta, igual que en
el notebook: se lee de la variable de entorno ``MTABOLIC_SAL_CODIGOS`` o, si no existe, del
archivo ``.sal_codigos`` de la carpeta de salidas (lo crea el notebook). La sal nunca se imprime.

Sale con 0 si no hay ninguna diferencia, 1 si hay alguna y 2 si faltan los libros, las
salidas del notebook o la sal, o ante un error no controlado (del que solo se imprime el tipo,
nunca el mensaje ni el traceback, que podrian citar datos de los libros). Solo imprime
conteos, cifras agregadas y la posicion de cada diferencia: nunca nombres, cedulas, celulares,
correos, codigos de socio, la sal ni rutas (RN-03.5, RT-13). No escribe archivos.
"""

import argparse
import hashlib
import hmac
import json
import math
import os
import sys
import warnings
from collections import defaultdict
from collections.abc import Iterator, Mapping, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, field, replace
from datetime import UTC, date, datetime
from decimal import ROUND_HALF_EVEN, Decimal
from pathlib import Path
from typing import TextIO

import openpyxl
import sklearn

from app.application.historial.ports import MembresiaParaPrediccion, NuevaMembresia
from app.application.historial.procesar_historico import ProcesarHistorico
from app.application.listas.dtos import GenerarListaDTO
from app.application.listas.generar_lista import GenerarLista
from app.application.listas.ports import (
    ItemLista,
    ListaDetalle,
    ListaResumen,
    NuevaLista,
    NuevaPrediccion,
)
from app.application.modelo.entrenar_modelo import EntrenarModeloInicial
from app.application.modelo.mapeo import ciclo_modelo_desde
from app.application.modelo.ports import EstadoSistema, ModeloGuardado, NuevoModelo
from app.domain.common.enums import EstadoMigracion
from app.domain.common.predictor import FeatureRow, Prediction, TrainedModel
from app.domain.historico.registro import RegistroHistorico
from app.domain.modelo.cartera import seleccionar_cartera, ya_renovo
from app.domain.modelo.variables import CicloModelo
from app.infrastructure.historico.lector_libros import LectorLibrosExcel
from app.infrastructure.ml.random_forest import RandomForestPredictor

FECHA_GENERACION = date(2026, 8, 31)
HORIZONTE_DIAS = 30
TOP_K = 25
CARPETA_SALIDAS = Path(__file__).resolve().parents[2] / "salidas"
ARCHIVO_PARAMETROS = Path("m01") / "despues" / "parametros_modelo_operativo.json"
ARCHIVO_LISTA = Path("mvp") / "lista_operativa_2026-09.xlsx"
HOJA_LISTA = "cartera_completa"
# Sal secreta de los codigos de socio: variable de entorno y, en su defecto, archivo en
# ``salidas/`` (git-ignorada) que genera el notebook la primera vez.
VARIABLE_SAL = "MTABOLIC_SAL_CODIGOS"
ARCHIVO_SAL = Path(".sal_codigos")
SEPARADOR_FACTORES = "; "
SIN_FACTORES = "Sin factores de riesgo destacados"
_OCHO_DECIMALES = Decimal("0.00000001")

# Las salidas del notebook M01 se generaron con scikit-learn 1.8.0 (la version fijada en la
# app); con otra version el bosque cambia y las probabilidades no coinciden.
SKLEARN_NOTEBOOK = "1.8.0"

# Claves de ``cartera_septiembre`` en el JSON del notebook (texto de la tabla 13.8).
_CLAVES_CARTERA: Mapping[str, str] = {
    "cartera": "Ciclos que vencen entre el 01-09-2026 y el 30-09-2026",
    "ya_renovaron": "ya registraron una membresía posterior (renovación anticipada)",
    "a_priorizar": "Ciclos a priorizar",
    "socios": "Socios distintos a priorizar",
}


class ReferenciaNoDisponibleError(Exception):
    """Faltan las salidas del notebook o no tienen la forma esperada."""


class SalNoDisponibleError(Exception):
    """No hay sal para los codigos de socio (ni variable de entorno ni archivo)."""


@dataclass(frozen=True, slots=True)
class FilaReferencia:
    """Una fila de la lista del notebook: codigo anonimo, probabilidad, nivel y factores."""

    codigo: str
    probabilidad: Decimal
    nivel: str
    factores: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class Referencia:
    """Salidas del notebook M01 con que se compara la app."""

    conteos: Mapping[str, int]
    variables: tuple[tuple[str, ...], tuple[str, ...]]
    medianas: Mapping[str, Decimal]
    referencia_factores: Mapping[str, float | str | None]
    lista: tuple[FilaReferencia, ...]


def codigo_anonimo(identidad: str, sal: bytes) -> str:
    """``id_anonimo`` del notebook: ``MTB-`` + 8 primeros hex en mayusculas del HMAC-SHA256.

    Sin la sal el codigo no se puede recalcular a partir de la cedula o el nombre.
    """
    huella = hmac.new(sal, identidad.encode("utf-8"), hashlib.sha256).hexdigest()
    return "MTB-" + huella[:8].upper()


def cargar_sal(carpeta: Path) -> bytes:
    """Sal de los codigos: ``MTABOLIC_SAL_CODIGOS`` o, si no existe, ``carpeta/.sal_codigos``."""
    valor = os.environ.get(VARIABLE_SAL, "").strip()
    if not valor:
        ruta = carpeta / ARCHIVO_SAL
        valor = ruta.read_text(encoding="utf-8").strip() if ruta.is_file() else ""
    if not valor:
        raise SalNoDisponibleError(
            f"No se encontro la sal de los codigos ({VARIABLE_SAL} o el archivo "
            f"{ARCHIVO_SAL} de las salidas del notebook)."
        )
    return valor.encode("utf-8")


def a_ocho_decimales(probabilidad: float | Decimal) -> Decimal:
    """Probabilidad con la precision que guarda la app (``NUMERIC(10,8)``)."""
    return Decimal(probabilidad).quantize(_OCHO_DECIMALES, rounding=ROUND_HALF_EVEN)


def separar_factores(texto: str) -> tuple[str, ...]:
    """Factores del notebook (unidos con ``"; "``) como la tupla que guarda la app."""
    if not texto or texto == SIN_FACTORES:
        return ()
    return tuple(texto.split(SEPARADOR_FACTORES))


def localizar_libros(carpeta: Path) -> list[Path]:
    """Libros ``*Control*Gym*.xlsx`` de la carpeta, ordenados; ignora bloqueos ``~$``."""
    return sorted(
        ruta
        for ruta in carpeta.glob("*.xlsx")
        if not ruta.name.startswith("~$")
        and "CONTROL" in ruta.name.upper()
        and "GYM" in ruta.name.upper()
    )


# ---------------------------------------------------------------- salidas del notebook


def _leer_lista(ruta: Path) -> tuple[FilaReferencia, ...]:
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        if HOJA_LISTA not in libro.sheetnames:
            raise ReferenciaNoDisponibleError("La lista del notebook no tiene la hoja esperada.")
        filas = libro[HOJA_LISTA].iter_rows(values_only=True)
        cabecera = [str(c) for c in next(filas)]
        try:
            i_codigo, i_prob, i_nivel, i_factores = (
                cabecera.index(c)
                for c in ("socio_anonimo", "probabilidad", "nivel_riesgo", "factores principales")
            )
        except ValueError as exc:
            raise ReferenciaNoDisponibleError(
                "La lista del notebook no tiene las columnas."
            ) from exc
        return tuple(
            FilaReferencia(
                codigo=str(fila[i_codigo]),
                probabilidad=a_ocho_decimales(float(str(fila[i_prob]))),
                nivel=str(fila[i_nivel]),
                factores=separar_factores(str(fila[i_factores] or "")),
            )
            for fila in filas
            if fila[i_codigo] is not None
        )
    finally:
        libro.close()


def cargar_referencia(carpeta: Path) -> Referencia:
    """Lee las salidas del notebook de ``carpeta`` (``salidas/``)."""
    parametros_ruta = carpeta / ARCHIVO_PARAMETROS
    lista_ruta = carpeta / ARCHIVO_LISTA
    if not parametros_ruta.is_file() or not lista_ruta.is_file():
        raise ReferenciaNoDisponibleError("No se encontraron las salidas del notebook M01.")
    parametros = json.loads(parametros_ruta.read_text(encoding="utf-8"))
    try:
        cartera = parametros["cartera_septiembre"]
        conteos = {nombre: int(cartera[clave]) for nombre, clave in _CLAVES_CARTERA.items()}
        conteos["ciclos_entrenamiento"] = int(parametros["ciclos_entrenamiento_operativo"])
        medianas_json = parametros["medianas_monto_operativo"]
        medianas = {plan: Decimal(str(v)) for plan, v in medianas_json["por_plan"].items()}
        medianas["general"] = Decimal(str(medianas_json["general"]))
        variables = (
            tuple(parametros["variables_numericas"]),
            tuple(parametros["variables_categoricas"]),
        )
        referencia_factores = dict(parametros["referencia_factores"])
    except (KeyError, TypeError, ValueError) as exc:
        raise ReferenciaNoDisponibleError(
            "Los parametros del notebook no tienen la forma."
        ) from exc
    return Referencia(
        conteos=conteos,
        variables=variables,
        medianas=medianas,
        referencia_factores=referencia_factores,
        lista=_leer_lista(lista_ruta),
    )


# ---------------------------------------------------------------- unidad de trabajo en memoria


class _RelojFijo:
    def __init__(self, instante: datetime) -> None:
        self._instante = instante

    def now(self) -> datetime:
        return self._instante


class _PredictorContador:
    """Delega en el adaptador real y cuenta las filas que ``GenerarLista`` manda a predecir."""

    def __init__(self, predictor: RandomForestPredictor) -> None:
        self._predictor = predictor
        self.filas_predichas = 0

    def train(
        self,
        features: Sequence[FeatureRow],
        labels: Sequence[int],
        numericas: Sequence[str],
        categoricas: Sequence[str],
        referencia: Sequence[FeatureRow] | None = None,
    ) -> TrainedModel:
        return self._predictor.train(features, labels, numericas, categoricas, referencia)

    def predict(self, artefacto: bytes, features: Sequence[FeatureRow]) -> list[Prediction]:
        self.filas_predichas += len(features)
        return self._predictor.predict(artefacto, features)


class _Membresias:
    def __init__(self, membresias: Sequence[MembresiaParaPrediccion]) -> None:
        self._membresias = list(membresias)

    def insertar_en_bloque(self, membresias: Sequence[NuevaMembresia]) -> int:
        raise NotImplementedError("La verificacion no inserta membresias.")

    def listar_para_prediccion(self) -> list[MembresiaParaPrediccion]:
        return list(self._membresias)


class _Modelos:
    def __init__(self, modelo: ModeloGuardado) -> None:
        self._modelo = modelo

    def crear(self, modelo: NuevoModelo) -> int:
        raise NotImplementedError("La verificacion no guarda modelos.")

    def obtener_vigente(self) -> ModeloGuardado | None:
        return self._modelo


class _Estado:
    def obtener(self) -> EstadoSistema:
        return EstadoSistema(EstadoMigracion.COMPLETADA, None)

    def fijar_migracion(self, estado: EstadoMigracion) -> None:
        raise NotImplementedError("La verificacion no cambia el estado.")

    def fijar_proxima_fecha(self, fecha: date) -> None:
        raise NotImplementedError("La verificacion no cambia el estado.")


class _Cargas:
    def ultima_fecha_carga_hasta(self, instante: datetime) -> datetime | None:
        return None


class _Listas:
    """Guarda la lista en memoria; los items no llevan nombre, celular ni cedula."""

    def __init__(
        self,
        modelo: ModeloGuardado,
        membresias: Sequence[MembresiaParaPrediccion],
        creada_en: datetime,
    ) -> None:
        self._modelo = modelo
        self._membresias = {m.membresia_id: m for m in membresias}
        self._creada_en = creada_en
        self.guardadas: list[tuple[NuevaLista, list[NuevaPrediccion]]] = []

    def crear(self, lista: NuevaLista, predicciones: Sequence[NuevaPrediccion]) -> int:
        self.guardadas.append((lista, list(predicciones)))
        return len(self.guardadas)

    def obtener(self, lista_id: int) -> ListaDetalle | None:
        if not 1 <= lista_id <= len(self.guardadas):
            return None
        lista, predicciones = self.guardadas[lista_id - 1]
        return ListaDetalle(
            id=lista_id,
            modelo_id=lista.modelo_id,
            modelo_version=self._modelo.version,
            fecha_generacion=lista.fecha_generacion,
            creada_en=self._creada_en,
            horizonte_dias=lista.horizonte_dias,
            top_k=lista.top_k,
            umbral_alto=self._modelo.umbral_alto,
            umbral_medio=self._modelo.umbral_medio,
            items=tuple(
                self._item(i, p)
                for i, p in enumerate(sorted(predicciones, key=lambda p: p.prioridad), 1)
            ),
            nombre=lista.nombre,
        )

    def _item(self, prediccion_id: int, p: NuevaPrediccion) -> ItemLista:
        m = self._membresias[p.membresia_id]
        if m.fecha_vencimiento is None:
            raise ValueError("Una membresia de la lista no tiene vencimiento.")
        return ItemLista(
            prediccion_id=prediccion_id,
            membresia_id=p.membresia_id,
            socio_id=p.socio_id,
            prioridad=p.prioridad,
            socio_nombre="",
            celular=None,
            cedula=None,
            plan=m.plan,
            fecha_vencimiento=m.fecha_vencimiento,
            monto_cobrado=m.monto_cobrado,
            score=p.score,
            nivel_riesgo=p.nivel_riesgo,
            factores_principales=list(p.factores_principales),
            dentro_top_k=p.dentro_top_k,
        )

    def ultima(self) -> ListaDetalle | None:
        return self.obtener(len(self.guardadas))

    def listar(self) -> list[ListaResumen]:
        return []


@dataclass(frozen=True, slots=True)
class _Contexto:
    listas: _Listas
    modelos: _Modelos
    estado: _Estado
    membresias: _Membresias
    cargas: _Cargas

    def commit(self) -> None:
        """Nada que confirmar: todo esta en memoria."""


# ---------------------------------------------------------------- ejecucion


@dataclass(frozen=True, slots=True)
class FilaApp:
    """Fila de la lista generada por la app, anonimizada."""

    prioridad: int
    codigo: str
    score: Decimal
    nivel: str
    factores: tuple[str, ...]
    dentro_top_k: bool


@dataclass(frozen=True, slots=True)
class ResultadoApp:
    conteos: dict[str, int]
    filas: tuple[FilaApp, ...]
    umbral_alto: Decimal
    umbral_medio: Decimal
    variables: tuple[tuple[str, ...], tuple[str, ...]] = ((), ())
    medianas: Mapping[str, Decimal] = field(default_factory=dict)
    referencia_factores: Mapping[str, float | str | None] = field(default_factory=dict)


def _membresias(registros: Sequence[RegistroHistorico]) -> list[MembresiaParaPrediccion]:
    """Membresias como las guarda la migracion: una por registro, ids en orden de lectura."""
    socio_de: dict[str, int] = {}
    for r in registros:
        socio_de.setdefault(r.identidad, len(socio_de) + 1)
    return [
        MembresiaParaPrediccion(
            membresia_id=i,
            socio_id=socio_de[r.identidad],
            plan=r.plan,
            forma_pago=r.forma_pago,
            fecha_inicio=r.fecha_inicio,
            fecha_vencimiento=r.fecha_vencimiento,
            fecha_actividad=r.fecha_actividad,
            monto_cobrado=r.monto_cobrado,
            inicio_aproximado=r.inicio_aproximado,
            # Decision (c) de R-005: todas "con cedula" para comparar con el notebook, que no
            # excluye a los socios sin cedula; la app si los excluye de la lista.
            socio_con_cedula=True,
        )
        for i, r in enumerate(registros, 1)
    ]


def _conteo_cartera(ciclos: Sequence[CicloModelo], fecha: date, horizonte: int) -> tuple[int, int]:
    """(ciclos en la cartera, ciclos de la cartera cuyo socio ya renovo)."""
    por_socio: defaultdict[str, list[CicloModelo]] = defaultdict(list)
    for c in ciclos:
        por_socio[c.clave_socio].append(c)
    cartera = seleccionar_cartera(ciclos, fecha, horizonte)
    return len(cartera), sum(ya_renovo(c, por_socio[c.clave_socio]) for c in cartera)


def _texto_json(metadata: Mapping[str, object], clave: str) -> object:
    valor = metadata.get(clave)
    return json.loads(valor) if isinstance(valor, str) else valor


def ejecutar_en_memoria(
    rutas: Sequence[Path], fecha: date, horizonte: int, top_k: int, sal: bytes
) -> ResultadoApp:
    """Historico -> modelo v1 -> lista de ``fecha``, todo en memoria y sin base de datos."""
    # openpyxl avisa de celdas con formato de fecha y serial invalido citando el valor de la
    # celda (puede ser una cedula mal ubicada): se silencian para no filtrar datos personales.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        procesado = ProcesarHistorico(LectorLibrosExcel()).ejecutar(rutas)

    predictor = RandomForestPredictor()
    entrenado = EntrenarModeloInicial(predictor).ejecutar(procesado.ciclos, procesado.fecha_corte)
    instante = datetime(fecha.year, fecha.month, fecha.day, tzinfo=UTC)
    modelo = ModeloGuardado(
        id=1,
        version=1,
        fecha_entrenamiento=instante,
        ciclos_entrenamiento=entrenado.ciclos_entrenamiento,
        periodo_desde=entrenado.periodo_desde,
        periodo_hasta=entrenado.periodo_hasta,
        prevalencia=entrenado.prevalencia,
        umbral_alto=entrenado.umbrales.alto,
        umbral_medio=entrenado.umbrales.medio,
        porcentajes_umbral=entrenado.porcentajes_umbral,
        artefacto=entrenado.artefacto,
        metadata_tecnica=entrenado.metadata_tecnica,
    )
    registros = procesado.registros
    membresias = _membresias(registros)
    listas = _Listas(modelo, membresias, instante)
    contexto = _Contexto(listas, _Modelos(modelo), _Estado(), _Membresias(membresias), _Cargas())

    @contextmanager
    def uow() -> Iterator[_Contexto]:
        yield contexto

    contador = _PredictorContador(predictor)
    # Decision (c) de R-005: la verificacion sigue su propio camino. El tope de 31 dias es una
    # validacion de entrada de la API; aqui se admite cualquier horizonte (>= 1) para comparar
    # con el notebook, por eso el DTO se construye sin esa validacion.
    if horizonte < 1 or top_k < 1:
        raise ValueError("El horizonte y K deben ser positivos.")
    parametros = GenerarListaDTO.model_construct(
        fecha_referencia=fecha, horizonte_dias=horizonte, top_k=top_k
    )
    GenerarLista(uow, contador, _RelojFijo(instante)).ejecutar(parametros)
    _, predicciones = listas.guardadas[-1]

    ciclos = [replace(ciclo_modelo_desde(c), ref=i) for i, c in enumerate(procesado.ciclos, 1)]
    en_cartera, ya_renovaron = _conteo_cartera(ciclos, fecha, horizonte)
    filas = tuple(
        FilaApp(
            prioridad=p.prioridad,
            codigo=codigo_anonimo(registros[p.membresia_id - 1].identidad, sal),
            score=p.score,
            nivel=p.nivel_riesgo.value,
            factores=p.factores_principales,
            dentro_top_k=p.dentro_top_k,
        )
        for p in sorted(predicciones, key=lambda p: p.prioridad)
    )
    medianas: dict[str, Decimal] = {
        plan.value: valor for plan, valor in entrenado.medianas.por_plan.items()
    }
    if entrenado.medianas.general is not None:
        medianas["general"] = entrenado.medianas.general
    metadata = entrenado.metadata_tecnica
    numericas = _texto_json(metadata, "variables_numericas")
    categoricas = _texto_json(metadata, "variables_categoricas")
    referencia = _texto_json(metadata, "referencia_factores")
    return ResultadoApp(
        conteos={
            "cartera": en_cartera,
            "ya_renovaron": ya_renovaron,
            "a_priorizar": contador.filas_predichas,
            "socios": len(filas),
            "ciclos_entrenamiento": entrenado.ciclos_entrenamiento,
        },
        filas=filas,
        umbral_alto=modelo.umbral_alto,
        umbral_medio=modelo.umbral_medio,
        variables=(
            tuple(numericas) if isinstance(numericas, list) else (),
            tuple(categoricas) if isinstance(categoricas, list) else (),
        ),
        medianas=medianas,
        referencia_factores=dict(referencia) if isinstance(referencia, dict) else {},
    )


# ---------------------------------------------------------------- conciliacion


@dataclass
class Conciliacion:
    """Lineas a imprimir y numero de diferencias por apartado."""

    lineas: list[str] = field(default_factory=list)
    diferencias: dict[str, int] = field(default_factory=dict)

    def contar(self, apartado: str, distintas: int) -> None:
        self.diferencias[apartado] = self.diferencias.get(apartado, 0) + distintas

    @property
    def total(self) -> int:
        return sum(self.diferencias.values())


def _conciliar_conteos(
    obtenidos: Mapping[str, int], referencia: Mapping[str, int], c: Conciliacion
) -> None:
    c.lineas.append(f"{'conteo':<22}  {'notebook':>8}  {'app':>8}  resultado")
    distintos = 0
    for nombre, esperado in referencia.items():
        obtenido = obtenidos.get(nombre)
        resultado = "ok" if obtenido == esperado else "DIFIERE"
        distintos += resultado == "DIFIERE"
        c.lineas.append(f"{nombre:<22}  {esperado:>8}  {obtenido!s:>8}  {resultado}")
    c.contar("conteos", distintos)


def _conciliar_modelo(app: ResultadoApp, ref: Referencia, c: Conciliacion) -> None:
    variables_ok = app.variables == ref.variables
    c.contar("variables", 0 if variables_ok else 1)
    c.lineas.append(
        f"variables del modelo ({len(ref.variables[0])} numericas, "
        f"{len(ref.variables[1])} categoricas): {'ok' if variables_ok else 'DIFIERE'}"
    )
    medianas_distintas = sorted(
        clave
        for clave in set(ref.medianas) | set(app.medianas)
        if ref.medianas.get(clave) != app.medianas.get(clave)
    )
    c.contar("medianas", len(medianas_distintas))
    c.lineas.append(
        f"medianas de monto ({len(ref.medianas)}): "
        + ("ok" if not medianas_distintas else "DIFIEREN " + ", ".join(medianas_distintas))
    )
    referencia_distinta = sorted(
        v
        for v in set(ref.referencia_factores) | set(app.referencia_factores)
        if not _mismo_valor(ref.referencia_factores.get(v), app.referencia_factores.get(v))
    )
    c.contar("referencia_factores", len(referencia_distinta))
    c.lineas.append(
        f"referencia de los factores ({len(ref.referencia_factores)} variables): "
        + ("ok" if not referencia_distinta else "DIFIERE " + ", ".join(referencia_distinta))
    )


def _mismo_valor(a: object, b: object) -> bool:
    if isinstance(a, int | float) and isinstance(b, int | float):
        return math.isclose(float(a), float(b), rel_tol=0, abs_tol=1e-9)
    return a == b


def _grupos(lista: Sequence[FilaReferencia]) -> list[tuple[int, int]]:
    """Rangos ``[inicio, fin)`` de posiciones consecutivas con la misma probabilidad."""
    grupos: list[tuple[int, int]] = []
    inicio = 0
    for i in range(1, len(lista) + 1):
        if i == len(lista) or lista[i].probabilidad != lista[inicio].probabilidad:
            grupos.append((inicio, i))
            inicio = i
    return grupos


def _conciliar_lista(
    filas: Sequence[FilaApp], lista: Sequence[FilaReferencia], top_k: int, c: Conciliacion
) -> None:
    por_codigo = {f.codigo: f for f in filas}
    prob = nivel = factores = ausentes = 0
    detalle: list[str] = []
    for posicion, ref in enumerate(lista, 1):
        app = por_codigo.get(ref.codigo)
        if app is None:
            ausentes += 1
            detalle.append(f"posicion {posicion}: el socio del notebook no esta en la lista")
            continue
        if app.score != ref.probabilidad:
            prob += 1
            detalle.append(
                f"posicion {posicion}: probabilidad notebook {ref.probabilidad} / app {app.score}"
            )
        if app.nivel != ref.nivel:
            nivel += 1
            detalle.append(f"posicion {posicion}: nivel notebook {ref.nivel} / app {app.nivel}")
        if app.factores != ref.factores:
            factores += 1
            detalle.append(
                f"posicion {posicion}: factores notebook {list(ref.factores)} / "
                f"app {list(app.factores)}"
            )
    sobrantes = len({f.codigo for f in filas} - {r.codigo for r in lista})

    orden = 0
    grupos = _grupos(lista)
    for inicio, fin in grupos:
        esperados = {r.codigo for r in lista[inicio:fin]}
        obtenidos = {f.codigo for f in filas[inicio:fin]}
        if esperados != obtenidos:
            orden += fin - inicio
            detalle.append(f"posiciones {inicio + 1}-{fin}: otro conjunto de socios en ese tramo")
    empates = [(i, f) for i, f in grupos if f - i > 1]
    en_el_corte = [(i, f) for i, f in empates if i < top_k < f]
    # Informativo: con un empate en el corte, cual de los empatados entra es arbitrario en el
    # notebook; fuera de ese caso una diferencia aqui ya la cuenta el orden.
    top_igual = {f.codigo for f in filas[:top_k]} == {r.codigo for r in lista[:top_k]}

    c.contar("filas", ausentes + sobrantes)
    c.contar("probabilidades", prob)
    c.contar("niveles", nivel)
    c.contar("factores", factores)
    c.contar("orden", orden)
    c.lineas.append(
        f"lista: notebook {len(lista)} filas | app {len(filas)} filas | ausentes {ausentes} | "
        f"sobrantes {sobrantes}"
    )
    c.lineas.append(
        f"probabilidades distintas (8 decimales): {prob} | niveles distintos: {nivel} | "
        f"factores distintos: {factores} | posiciones fuera de orden: {orden}"
    )
    c.lineas.append(
        f"grupos de empate (misma probabilidad): {len(empates)}, con "
        f"{sum(f - i for i, f in empates)} socios; cruzan el corte del top {top_k}: "
        f"{len(en_el_corte)} (el top {top_k} se elige dentro del empate)"
    )
    c.lineas.append(
        f"socios dentro del top {top_k}: "
        + ("los mismos que en el notebook" if top_igual else "distintos (informativo)")
    )
    c.lineas.extend(detalle)


def _niveles(filas: Sequence[FilaApp]) -> str:
    return ", ".join(f"{n} {sum(f.nivel == n for f in filas)}" for n in ("ALTO", "MEDIO", "BAJO"))


def main(
    argv: Sequence[str] | None = None,
    *,
    referencia: Referencia | None = None,
    sal: bytes | None = None,
    salida: TextIO = sys.stdout,
) -> int:
    try:
        return _verificar(argv, referencia, sal, salida)
    except Exception as exc:  # ultima barrera: ni mensaje ni traceback
        # El mensaje o el traceback podrian citar valores de los libros (cedulas, nombres) o
        # rutas locales: solo se imprime el tipo del error.
        print(f"error no controlado: {type(exc).__name__}", file=salida)
        return 2


def _verificar(
    argv: Sequence[str] | None,
    referencia: Referencia | None,
    sal: bytes | None,
    salida: TextIO,
) -> int:
    parser = argparse.ArgumentParser(description="Verifica la lista frente al notebook M01.")
    parser.add_argument("carpeta", type=Path, help="carpeta con los libros Control Gym")
    parser.add_argument("--referencia", type=Path, default=CARPETA_SALIDAS)
    parser.add_argument("--fecha", type=date.fromisoformat, default=FECHA_GENERACION)
    parser.add_argument("--horizonte", type=int, default=HORIZONTE_DIAS)
    parser.add_argument("--top-k", type=int, default=TOP_K)
    argumentos = parser.parse_args(argv)

    def imprimir(texto: str = "") -> None:
        print(texto, file=salida)

    rutas = localizar_libros(argumentos.carpeta)
    if not rutas:
        imprimir("No se encontraron libros Control Gym en la carpeta indicada.")
        return 2
    if referencia is None:
        try:
            referencia = cargar_referencia(argumentos.referencia)
        except ReferenciaNoDisponibleError as exc:
            imprimir(str(exc))
            return 2
    if sal is None:
        try:
            sal = cargar_sal(argumentos.referencia)
        except SalNoDisponibleError as exc:
            imprimir(str(exc))
            return 2

    resultado = ejecutar_en_memoria(
        rutas, argumentos.fecha, argumentos.horizonte, argumentos.top_k, sal
    )
    conciliacion = Conciliacion()
    _conciliar_conteos(resultado.conteos, referencia.conteos, conciliacion)
    _conciliar_modelo(resultado, referencia, conciliacion)
    _conciliar_lista(resultado.filas, referencia.lista, argumentos.top_k, conciliacion)

    imprimir(
        f"lista de {argumentos.fecha.isoformat()} | horizonte {argumentos.horizonte} dias | "
        f"top {argumentos.top_k} | libros: {len(rutas)}"
    )
    for linea in conciliacion.lineas:
        imprimir(linea)
    imprimir(
        f"umbrales de la app: alto {resultado.umbral_alto} | medio {resultado.umbral_medio} "
        f"(fijos del notebook); niveles en la lista: {_niveles(resultado.filas)}"
    )
    imprimir(
        "diferencias: "
        + ", ".join(f"{k} {v}" for k, v in conciliacion.diferencias.items())
        + f" | total {conciliacion.total}"
    )
    if sklearn.__version__ != SKLEARN_NOTEBOOK:
        imprimir(
            f"aviso: scikit-learn {sklearn.__version__} instalado; las salidas del notebook se "
            f"generaron con {SKLEARN_NOTEBOOK} (el bosque cambia entre versiones)"
        )
    return 0 if conciliacion.total == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
