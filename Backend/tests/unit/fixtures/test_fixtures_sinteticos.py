"""Estructura y conteos de los fixtures sintéticos (D2-OP-01).

Recalcula desde los ``.xlsx`` (con openpyxl y un oráculo mínimo de las reglas del notebook,
secciones 2.4 a 5.5) los conteos de ``esperado.py``, y comprueba que los fixtures no vienen de
los libros reales. El oráculo vive solo aquí: no se importa nada de ``app``.
"""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.fixtures import esperado
from tests.fixtures import generar_fixtures as gen

pytestmark = pytest.mark.unit

BACKEND = Path(__file__).resolve().parents[3]
DIR_FIXTURES = BACKEND / "tests" / "fixtures"
DIR_REALES = BACKEND.parent / "plantillas"
LIBROS = (esperado.LIBRO_2025, esperado.LIBRO_2026)
PLANTILLAS = tuple(esperado.FILAS_CON_DATOS)
ARCHIVOS = tuple(f"historico/{n}" for n in LIBROS) + tuple(f"plantillas/{n}" for n in PLANTILLAS)

# ------------------------------------------------------------------ oráculo (notebook)

MESES = [
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE",
]
ALIAS = {
    "nombre": ("NOMBRES Y APELLIDOS", "NOMBRES Y APELLIDOS DEL SOCIO"),
    "cedula": ("CEDULA",),
    "plan": ("TIPO DE PLAN",),
    "inicio": ("FECHA INSCRIPCION", "FECHA INICIO"),
    "fin": ("FECHA FIN PLAN", "FECHA FIN"),
    "monto": ("MONTO A COBRAR", "MONTO COBRAR"),
}
PAGOS_EN_PLAN = ("DATAFAST", "TRANSF", "EFECTIVO", "TARJETA", "DEPOSITO")


def _sin_tildes(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()


def _encabezado(valor: object) -> str:
    return "" if valor is None else re.sub(r"\s+", " ", _sin_tildes(str(valor).strip().upper()))


def _fecha(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    numerico = isinstance(valor, int | float) and not isinstance(valor, bool)
    if numerico and 20000 <= float(str(valor)) <= 60000:
        return date(1899, 12, 30) + timedelta(days=int(float(str(valor))))
    texto = "" if valor is None else str(valor).strip()
    partes = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", texto)
    return date(int(partes[3]), int(partes[2]), int(partes[1])) if partes else None


def _monto(valor: object) -> Decimal | None:
    if valor is None:
        return None
    hallado = re.search(r"-?\d+(?:\.\d+)?", str(valor).replace(",", "."))
    return Decimal(hallado.group(0)) if hallado else None


def _nombre(valor: object) -> str:
    texto = _sin_tildes("" if valor is None else str(valor)).upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z ]", " ", texto)).strip()


def _cedula(valor: object) -> str | None:
    digitos = re.sub(r"\D", "", "" if valor is None else str(valor))
    return digitos if len(digitos) >= 9 else None


def _plan(texto: object, duracion: int | None) -> str:
    plan = str(texto or "").upper().strip()
    if any(p in plan for p in PAGOS_EN_PLAN):
        plan = ""
    reglas = (
        ("DIARIO", ("DIARIO",)),
        ("ANUAL", ("ANUAL", "AÑO", "ANO")),
        ("SEMESTRAL", ("SEMESTR", "6 MES")),
        ("TRIMESTRAL", ("3 MES", "TRES MES", "TRIMEST")),
    )
    for categoria, claves in reglas:
        if any(c in plan for c in claves):
            return categoria
    if "MENSUAL" in plan or re.search(r"\b1\s*MES\b", plan) or "PLAN MES" in plan:
        return "MENSUAL"
    rangos = (("MENSUAL", 21, 45), ("TRIMESTRAL", 80, 130), ("SEMESTRAL", 160, 200))
    for categoria, minimo, maximo in (*rangos, ("ANUAL", 330, 400)):
        if duracion is not None and minimo <= duracion <= maximo:
            return categoria
    return "OTRO"


@dataclass
class Registro:
    anio: int
    hoja: str
    fila: int
    identidad: str
    por_cedula: bool
    plan: str
    fin: date | None
    duracion: int | None
    actividad: date
    aproximado: bool
    monto: Decimal | None


@dataclass
class Lectura:
    incluidas: dict[str, list[str]]
    omitidas: dict[str, list[str]]
    encabezados: dict[tuple[int, str], int]
    leidas: dict[tuple[int, str], int]
    sin_inicio: list[tuple[int, str]]
    crudas: list[tuple[int, str, int, date, dict[str, object]]]


def _leer(directorio: Path) -> Lectura:
    lectura = Lectura({}, {}, {}, {}, [], [])
    for archivo in LIBROS:
        anio = esperado.ANIO_POR_LIBRO[archivo]
        libro = load_workbook(directorio / archivo, data_only=True)
        lectura.incluidas[archivo], lectura.omitidas[archivo] = [], []
        for hoja in libro.worksheets:
            clave = hoja.title.strip().upper()
            if clave not in MESES:
                lectura.omitidas[archivo].append(hoja.title)
                continue
            lectura.incluidas[archivo].append(hoja.title)
            filas = [list(f) for f in hoja.iter_rows(values_only=True)]
            while filas and all(v is None for v in filas[-1]):
                filas.pop()
            cabecera = next(
                i
                for i, f in enumerate(filas[:10])
                if "CEDULA" in (t := " | ".join(_encabezado(v) for v in f)) and "NOMBRE" in t
            )
            titulos = [_encabezado(v) for v in filas[cabecera]]
            indices = {
                campo: next((i for i, t in enumerate(titulos) if any(a in t for a in alias)), None)
                for campo, alias in ALIAS.items()
            }
            lectura.encabezados[(anio, hoja.title)] = cabecera
            lectura.leidas[(anio, hoja.title)] = len(filas) - cabecera - 1
            if indices["inicio"] is None:
                lectura.sin_inicio.append((anio, hoja.title))
            mes = date(anio, MESES.index(clave) + 1, 1)
            for posicion in range(cabecera + 1, len(filas)):
                valores = {
                    c: (None if i is None else filas[posicion][i]) for c, i in indices.items()
                }
                lectura.crudas.append((anio, hoja.title, posicion + 1, mes, valores))
    return lectura


def _limpiar(lectura: Lectura) -> tuple[int, int, list[Registro]]:
    no_utiles, sin_identidad, registros = 0, 0, []
    for anio, hoja, fila, mes, v in lectura.crudas:
        crudo = "" if v["nombre"] is None else str(v["nombre"]).strip()
        mayus = crudo.upper()
        repetido = mayus in ("STATUS", "NOMBRES Y APELLIDOS", "NOMBRES Y APELLIDOS DEL SOCIO")
        if not crudo or re.search("TOTAL|SUBTOTAL|SUMA|MONTO", mayus) or repetido:
            no_utiles += 1
            continue
        nombre = _nombre(crudo)
        if len(nombre) < 5:
            sin_identidad += 1
            continue
        cedula = _cedula(v["cedula"])
        inicio, fin = _fecha(v["inicio"]), _fecha(v["fin"])
        duracion = (fin - inicio).days if inicio and fin else None
        registros.append(
            Registro(
                anio,
                hoja,
                fila,
                f"C:{cedula}" if cedula else f"N:{nombre}",
                cedula is not None,
                _plan(v["plan"], duracion),
                fin,
                duracion,
                inicio or mes,
                inicio is None,
                _monto(v["monto"]),
            )
        )
    return no_utiles, sin_identidad, registros


def _motivo(r: Registro, corte: date) -> str:
    if r.fin is None:
        return "SIN_VENCIMIENTO"
    if r.fin < date(2024, 12, 1):
        return "FECHA_FUERA_DE_RANGO"
    if r.fin > corte:
        return "VENCIMIENTO_POSTERIOR_AL_CORTE"
    if r.plan == "DIARIO":
        return "PASE_DIARIO"
    if r.duracion is not None and (r.duracion <= 21 or r.duracion > 400):
        return "DURACION_IMPLAUSIBLE"
    if r.fin + timedelta(days=30) > corte:
        return "VENTANA_INCOMPLETA"
    return "EVALUABLE"


def _renovo(r: Registro, registros: list[Registro]) -> bool:
    assert r.fin is not None
    return any(
        o is not r
        and o.identidad == r.identidad
        and o.plan != "DIARIO"
        and r.fin <= o.actividad <= r.fin + timedelta(days=30)
        for o in registros
    )


# ------------------------------------------------------------------ fixtures de pytest


@pytest.fixture(scope="module")
def generados(tmp_path_factory: pytest.TempPathFactory) -> Path:
    destino = tmp_path_factory.mktemp("fixtures")
    gen.generar(destino)
    return destino


@pytest.fixture(scope="module", params=["versionados", "regenerados"])
def directorio(request: pytest.FixtureRequest, generados: Path) -> Path:
    return DIR_FIXTURES if request.param == "versionados" else generados


# ------------------------------------------------------------------ libros históricos


def test_regenerar_produce_los_siete_archivos_con_los_mismos_bytes(tmp_path: Path) -> None:
    primera = gen.generar(tmp_path / "a")
    segunda = gen.generar(tmp_path / "b")

    relativas = sorted(p.relative_to(tmp_path / "a").as_posix() for p in primera)
    assert relativas == sorted(ARCHIVOS)
    for a, b in zip(primera, segunda, strict=True):
        assert a.read_bytes() == b.read_bytes(), a.name


def test_los_archivos_versionados_existen() -> None:
    for relativo in ARCHIVOS:
        assert (DIR_FIXTURES / relativo).is_file(), relativo


def test_hojas_encabezados_y_filas_leidas_coinciden(directorio: Path) -> None:
    lectura = _leer(directorio / "historico")

    for archivo in LIBROS:
        assert tuple(lectura.incluidas[archivo]) == esperado.HOJAS_INCLUIDAS[archivo]
        assert tuple(lectura.omitidas[archivo]) == esperado.HOJAS_OMITIDAS[archivo]
    assert lectura.encabezados == esperado.FILA_ENCABEZADO
    assert set(lectura.encabezados.values()) >= {0, 2, 4}
    assert lectura.leidas == esperado.FILAS_LEIDAS_POR_HOJA
    assert sum(lectura.leidas.values()) == esperado.FILAS_LEIDAS
    assert tuple(lectura.sin_inicio) == esperado.HOJAS_SIN_FECHA_INICIO


def test_limpieza_e_identidad_coinciden(directorio: Path) -> None:
    no_utiles, sin_identidad, registros = _limpiar(_leer(directorio / "historico"))
    socios = {r.identidad for r in registros}

    assert no_utiles == esperado.DESCARTADAS_NO_UTILES
    assert len(registros) + sin_identidad == esperado.UTILES
    assert sin_identidad == esperado.SIN_IDENTIDAD
    assert len(registros) == esperado.CON_IDENTIDAD
    assert len(socios) == esperado.SOCIOS_DISTINTOS
    assert sum(s.startswith("C:") for s in socios) == esperado.SOCIOS_POR_CEDULA
    assert sum(s.startswith("N:") for s in socios) == esperado.SOCIOS_POR_NOMBRE
    assert sum(r.por_cedula for r in registros) == esperado.REGISTROS_POR_CEDULA
    assert sum(not r.por_cedula for r in registros) == esperado.REGISTROS_POR_NOMBRE
    assert sum(r.aproximado for r in registros) == esperado.INICIO_APROXIMADO
    negativos = sum(r.monto is not None and r.monto < 0 for r in registros)
    assert negativos == esperado.MONTOS_INVALIDOS
    claves = Counter((r.identidad, r.actividad, r.fin, r.plan, r.monto) for r in registros)
    assert sum(n for n in claves.values() if n > 1) == esperado.DUPLICADOS_EXACTOS


def test_ciclos_etiquetas_y_caso_del_tutor_coinciden(directorio: Path) -> None:
    lectura = _leer(directorio / "historico")
    _, _, registros = _limpiar(lectura)
    ultimo_mes = max(mes for *_, mes, _ in lectura.crudas)
    siguiente = (ultimo_mes.replace(day=28) + timedelta(days=4)).replace(day=1)
    corte = siguiente - timedelta(days=1)
    motivos = {id(r): _motivo(r, corte) for r in registros}
    evaluables = [r for r in registros if motivos[id(r)] == "EVALUABLE"]
    renovaron = [r for r in evaluables if _renovo(r, registros)]
    vencimientos = [r.fin for r in registros if r.fin is not None]

    assert corte == esperado.FECHA_CORTE
    assert Counter(motivos.values()) == esperado.CICLOS_POR_MOTIVO
    estado_de = {"EVALUABLE": "EVALUABLE", "VENTANA_INCOMPLETA": "PENDIENTE"}
    estados = Counter(estado_de.get(m, "EXCLUIDO") for m in motivos.values())
    assert estados == esperado.CICLOS_POR_ESTADO
    assert len(evaluables) == esperado.EVALUABLES
    assert len(renovaron) == esperado.RENOVO
    assert len(evaluables) - len(renovaron) == esperado.NO_RENOVO
    assert Decimal(esperado.NO_RENOVO) / Decimal(esperado.EVALUABLES) == esperado.PREVALENCIA
    assert tuple((r.anio, r.hoja, r.fila) for r in renovaron) == esperado.RENOVACIONES
    assert (min(vencimientos), max(vencimientos)) == esperado.RANGO_VENCIMIENTOS
    fines_eval = [r.fin for r in evaluables if r.fin is not None]
    rango_eval = (min(fines_eval), max(fines_eval))
    assert rango_eval == esperado.RANGO_VENCIMIENTOS_EVALUABLES

    tutor = esperado.CASO_TUTOR
    [caso] = [
        r
        for r in registros
        if (r.anio, r.hoja, r.fila) == (tutor["anio"], tutor["hoja"], tutor["fila"])
    ]
    assert caso.fin == tutor["fecha_vencimiento"]
    assert motivos[id(caso)] == tutor["motivo"]


def test_los_libros_cubren_los_formatos_de_celda_pedidos(directorio: Path) -> None:
    fechas, montos, cedulas, planes = set(), set(), set(), set()
    for _, _, _, _, v in _leer(directorio / "historico").crudas:
        for campo in ("inicio", "fin"):
            if v[campo] is not None:
                fechas.add(type(v[campo]).__name__)
        if isinstance(v["monto"], str):
            montos.add(v["monto"])
        if isinstance(v["cedula"], str) and "-" in v["cedula"]:
            cedulas.add("guion")
        planes.add(str(v["plan"]).upper())

    assert {"datetime", "int", "str"} <= fechas  # fecha real, serial Excel y texto dd/mm/yyyy
    assert "39,50 USD" in montos
    assert cedulas == {"guion"}
    assert any(p in plan for plan in planes for p in PAGOS_EN_PLAN)
    assert "PASE DIARIO" in planes


# ------------------------------------------------------------------ plantillas


def _digito_valido(cedula: str) -> bool:
    if not re.fullmatch(r"\d{10}", cedula):
        return False
    total = 0
    for posicion, caracter in enumerate(cedula[:9]):
        producto = int(caracter) * (2 if posicion % 2 == 0 else 1)
        total += producto - 9 if producto > 9 else producto
    return (10 - total % 10) % 10 == int(cedula[9])


def _errores_fila(valores: list[object]) -> list[str]:
    errores = []
    if not _digito_valido(str(valores[2])):
        errores.append("cedula")
    if valores[5] not in gen.PLANES:
        errores.append("tipo_plan")
    inicio, fin = valores[6], valores[7]
    if not (isinstance(inicio, datetime) and isinstance(fin, datetime) and fin > inicio):
        errores.append("fecha_vencimiento")
    if valores[8] not in gen.FORMAS_PAGO or valores[0] not in gen.ESTADOS_REGISTRO:
        errores.append("catalogo")
    if len(_nombre(valores[1])) < 5 or not isinstance(valores[11], float | int):
        errores.append("obligatorio")
    return errores


@pytest.mark.parametrize("archivo", PLANTILLAS)
def test_plantillas_estructura_y_errores(directorio: Path, archivo: str) -> None:
    libro = load_workbook(directorio / "plantillas" / archivo)
    registros = libro["Registros"]

    assert tuple(libro.sheetnames) == esperado.HOJAS_PLANTILLA
    assert tuple(c.value for c in registros[1]) == gen.TITULOS_REGISTROS
    filas = [list(f[:14]) for f in registros.iter_rows(min_row=2, values_only=True)]
    con_datos = [(n, f) for n, f in enumerate(filas, start=2) if any(v is not None for v in f)]
    assert len(con_datos) == esperado.FILAS_CON_DATOS[archivo]
    assert [n for n, _ in con_datos] == list(range(2, len(con_datos) + 2))
    errores = {n: e for n, f in con_datos if (e := _errores_fila(f))}
    assert errores == {n: [c] for n, c in esperado.FILAS_CON_ERROR[archivo].items()}


def test_catalogos_de_la_plantilla_en_las_celdas_de_las_listas(directorio: Path) -> None:
    catalogos = load_workbook(directorio / "plantillas" / esperado.PLANTILLA_VALIDA)["Catalogos"]

    def rango(desde: int, hasta: int) -> tuple[object, ...]:
        return tuple(catalogos.cell(row=r, column=1).value for r in range(desde, hasta + 1))

    assert rango(5, 10) == gen.PLANES
    assert rango(14, 17) == gen.FORMAS_PAGO
    assert rango(21, 25) == gen.ESTADOS_REGISTRO


# ------------------------------------------------------------------ sin datos reales


def test_nombres_de_socios_salen_de_la_lista_ficticia(directorio: Path) -> None:
    ficticios = set(gen.NOMBRES_FICTICIOS)
    for *_, v in _leer(directorio / "historico").crudas:
        nombre = v["nombre"]
        if isinstance(nombre, str) and nombre.strip():
            etiqueta = re.search("TOTAL|SUBTOTAL|MONTO|NOMBRES Y APELLIDOS", nombre.upper())
            assert etiqueta or nombre in ficticios, "nombre fuera de la lista ficticia"
    for archivo in PLANTILLAS:
        registros = load_workbook(directorio / "plantillas" / archivo)["Registros"]
        for (nombre,) in registros.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
            assert nombre is None or nombre in ficticios


def test_ningun_fixture_viene_de_plantillas_reales() -> None:
    nombres_fixture = {Path(a).name for a in ARCHIVOS}
    assert all("sintetico" in n or n.startswith("plantilla_") for n in nombres_fixture)
    if DIR_REALES.is_dir():
        reales = {p.name for p in DIR_REALES.glob("*.xlsx")}
        assert not nombres_fixture & reales
        huellas_reales = {hashlib.sha256(p.read_bytes()).digest() for p in DIR_REALES.glob("*")}
        for relativo in ARCHIVOS:
            huella = hashlib.sha256((DIR_FIXTURES / relativo).read_bytes()).digest()
            assert huella not in huellas_reales, relativo


def test_git_no_ignora_los_fixtures() -> None:
    git = shutil.which("git")
    if git is None:
        pytest.skip("git no está disponible en este entorno")
    rutas = [str(DIR_FIXTURES / a) for a in ARCHIVOS]
    resultado = subprocess.run(  # noqa: S603 - argumentos fijos, sin shell
        [git, "check-ignore", "--no-index", *rutas],
        cwd=BACKEND,
        capture_output=True,
        text=True,
        check=False,
    )
    if resultado.returncode == 128:
        pytest.skip("Backend no está dentro de un repositorio git")
    assert resultado.returncode == 1, "git ignora algún fixture"
    assert resultado.stdout == ""
