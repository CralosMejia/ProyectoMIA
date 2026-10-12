"""Lector de la plantilla oficial de carga (R-007 T4, D4-BE-01).

Solo fixtures sinteticos (``tests/fixtures``) y libros construidos en ``tmp_path`` con
nombres y cedulas ficticios. Sin login (RT-12): no hay pruebas de aislamiento entre usuarios.
"""

import io
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]  # openpyxl no publica stubs

from app.application.cargas.ports import LectorPlantilla
from app.domain.cargas.errores import MotivoPlantillaInvalida, PlantillaInvalidaError
from app.domain.cargas.fila import COLUMNAS
from app.infrastructure.cargas import lector_plantilla
from app.infrastructure.cargas.lector_plantilla import (
    MAX_FILAS_ITERADAS,
    MAX_FILAS_PLANTILLA,
    TITULOS_PLANTILLA,
    LectorPlantillaOpenpyxl,
)
from tests.fixtures.esperado import FILAS_CON_DATOS, PLANTILLA_501_FILAS, PLANTILLA_VALIDA
from tests.fixtures.generar_fixtures import (
    TITULOS_REGISTROS,
    fila_plantilla,
    filas_plantillas,
    generar_plantilla,
)

pytestmark = pytest.mark.unit

FIXTURES = Path(__file__).resolve().parents[3] / "fixtures"
PLANTILLAS = FIXTURES / "plantillas"
LIBRO_ANTIGUO = FIXTURES / "historico" / "Control Gym 2025 sintetico.xlsx"


@pytest.fixture
def lector() -> LectorPlantillaOpenpyxl:
    return LectorPlantillaOpenpyxl()


def libro_en_bytes(hojas: dict[str, list[list[object]]]) -> bytes:
    """Libro minimo con las hojas y filas dadas (fila 1 = primera lista)."""
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for numero, valores in enumerate(filas, start=1):
            for columna, valor in enumerate(valores, start=1):
                if valor is not None:
                    hoja.cell(row=numero, column=columna, value=valor)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def motivo_de(lector: LectorPlantillaOpenpyxl, contenido: bytes) -> PlantillaInvalidaError:
    with pytest.raises(PlantillaInvalidaError) as capturado:
        lector.leer(contenido)
    return capturado.value


def test_lector_implements_the_port(lector: LectorPlantillaOpenpyxl) -> None:
    assert isinstance(lector, LectorPlantilla)


def test_official_titles_match_fixture_titles_a_to_n() -> None:
    assert TITULOS_REGISTROS[:14] == TITULOS_PLANTILLA
    assert len(TITULOS_PLANTILLA) == len(COLUMNAS)


def test_leer_valid_template_returns_six_rows_with_real_numbers_and_raw_values(
    lector: LectorPlantillaOpenpyxl,
) -> None:
    plantilla = lector.leer((PLANTILLAS / PLANTILLA_VALIDA).read_bytes())
    esperadas = filas_plantillas()[PLANTILLA_VALIDA]

    assert [f.fila for f in plantilla.filas] == [2, 3, 4, 5, 6, 7]
    for fila, valores in zip(plantilla.filas, esperadas, strict=True):
        crudos = [fila.valor(letra) for letra, _ in COLUMNAS]
        assert crudos[1:6] == valores[1:6]
        assert isinstance(fila.fecha_inicio, date | datetime)
        assert isinstance(fila.fecha_vencimiento, date | datetime)
        inicio = fila.fecha_inicio
        esperado = valores[6]
        assert isinstance(esperado, date)
        assert (inicio.date() if isinstance(inicio, datetime) else inicio) == esperado
        assert fila.asistencias == valores[12]


@pytest.mark.parametrize(
    "archivo", [a for a in FILAS_CON_DATOS if a != PLANTILLA_501_FILAS], ids=str
)
def test_leer_each_fixture_returns_expected_rows_with_data(
    lector: LectorPlantillaOpenpyxl, archivo: str
) -> None:
    plantilla = lector.leer((PLANTILLAS / archivo).read_bytes())

    assert len(plantilla.filas) == FILAS_CON_DATOS[archivo]


def test_leer_501_rows_raises_too_many_rows(lector: LectorPlantillaOpenpyxl) -> None:
    error = motivo_de(lector, (PLANTILLAS / PLANTILLA_501_FILAS).read_bytes())

    assert error.motivo is MotivoPlantillaInvalida.DEMASIADAS_FILAS
    assert str(error) == (
        "El archivo tiene más de 500 filas con datos; divídalo en varios archivos."
    )


def test_leer_exactly_500_rows_is_accepted(lector: LectorPlantillaOpenpyxl, tmp_path: Path) -> None:
    ruta = generar_plantilla(
        [fila_plantilla(i) for i in range(MAX_FILAS_PLANTILLA)], tmp_path / "p.xlsx"
    )

    plantilla = lector.leer(ruta.read_bytes())

    assert len(plantilla.filas) == 500
    assert plantilla.filas[-1].fila == 501


def test_leer_old_monthly_format_raises_old_format(lector: LectorPlantillaOpenpyxl) -> None:
    error = motivo_de(lector, LIBRO_ANTIGUO.read_bytes())

    assert error.motivo is MotivoPlantillaInvalida.FORMATO_ANTIGUO
    assert str(error) == (
        "El archivo tiene el formato antiguo de hojas mensuales; use la plantilla oficial de carga."
    )


def test_leer_without_registros_nor_months_raises_missing_sheet(
    lector: LectorPlantillaOpenpyxl,
) -> None:
    contenido = libro_en_bytes({"Hoja1": [list(TITULOS_PLANTILLA)]})

    error = motivo_de(lector, contenido)

    assert error.motivo is MotivoPlantillaInvalida.SIN_REGISTROS
    assert str(error) == "El archivo no tiene la hoja Registros."


@pytest.mark.parametrize("columna", [0, 2, 13])
def test_leer_changed_title_raises_titles_mismatch(
    lector: LectorPlantillaOpenpyxl, columna: int
) -> None:
    titulos: list[object] = list(TITULOS_PLANTILLA)
    titulos[columna] = "OTRO TITULO"
    contenido = libro_en_bytes({"Registros": [titulos, fila_plantilla(0)]})

    error = motivo_de(lector, contenido)

    assert error.motivo is MotivoPlantillaInvalida.TITULOS
    assert str(error) == "Los títulos de la hoja Registros no coinciden con la plantilla oficial."


def test_leer_missing_title_raises_titles_mismatch(lector: LectorPlantillaOpenpyxl) -> None:
    contenido = libro_en_bytes({"Registros": [list(TITULOS_PLANTILLA[:13])]})

    assert motivo_de(lector, contenido).motivo is MotivoPlantillaInvalida.TITULOS


def test_leer_titles_with_spaces_and_other_o_p_titles_are_accepted(
    lector: LectorPlantillaOpenpyxl,
) -> None:
    titulos: list[object] = [*(f"  {t} " for t in TITULOS_PLANTILLA), "X", "Y"]
    contenido = libro_en_bytes({"Registros": [titulos, fila_plantilla(0)]})

    assert len(lector.leer(contenido).filas) == 1


@pytest.mark.parametrize(
    "contenido",
    [b"", b"esto no es un libro", b"PK\x03\x04roto"],
    ids=["vacio", "texto", "zip-roto"],
)
def test_leer_non_xlsx_bytes_raises_damaged(
    lector: LectorPlantillaOpenpyxl, contenido: bytes
) -> None:
    error = motivo_de(lector, contenido)

    assert error.motivo is MotivoPlantillaInvalida.DANADO
    assert str(error) == "El archivo está dañado o no es un .xlsx válido."


def test_leer_zip_that_is_not_a_workbook_raises_damaged(lector: LectorPlantillaOpenpyxl) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archivo:
        archivo.writestr("hola.txt", "no es un libro")

    assert motivo_de(lector, buffer.getvalue()).motivo is MotivoPlantillaInvalida.DANADO


def test_leer_skips_empty_rows_keeping_real_row_number(lector: LectorPlantillaOpenpyxl) -> None:
    vacia: list[object] = [None] * 14
    solo_espacios: list[object] = ["  ", None, " "]
    contenido = libro_en_bytes(
        {
            "Registros": [
                list(TITULOS_PLANTILLA),
                fila_plantilla(0),
                vacia,
                solo_espacios,
                vacia,
                fila_plantilla(1),
            ]
        }
    )

    plantilla = lector.leer(contenido)

    assert [f.fila for f in plantilla.filas] == [2, 6]


def test_leer_ignores_columns_o_and_p(lector: LectorPlantillaOpenpyxl) -> None:
    solo_o_p: list[object] = [None] * 14 + ["=H3-G3", "REVISAR"]
    contenido = libro_en_bytes(
        {"Registros": [list(TITULOS_PLANTILLA), [*fila_plantilla(0), 30, "OK"], solo_o_p]}
    )

    plantilla = lector.leer(contenido)

    assert [f.fila for f in plantilla.filas] == [2]


def test_error_messages_never_contain_cell_values(lector: LectorPlantillaOpenpyxl) -> None:
    titulos: list[object] = list(TITULOS_PLANTILLA)
    titulos[1] = "Nombre Ficticio 1700000019"
    errores = [
        motivo_de(lector, libro_en_bytes({"Registros": [titulos, fila_plantilla(0)]})),
        motivo_de(lector, (PLANTILLAS / PLANTILLA_501_FILAS).read_bytes()),
    ]
    valores = [str(v) for v in fila_plantilla(0)[1:5]] + ["Nombre Ficticio", "1700000019"]

    for error in errores:
        assert not any(valor in str(error) for valor in valores)


# ---------------------------------------------------- S3 (R-013): zip bomb y filas iteradas

MENSAJE_DEMASIADO_GRANDE = (
    "El contenido del archivo es demasiado grande para una plantilla de carga; "
    "use la plantilla oficial."
)


def _con_entrada_extra(contenido: bytes, nombre: str, datos: bytes) -> bytes:
    """Copia el libro y anade una entrada comprimida (deflate) al zip."""
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(contenido)) as origen,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as destino,
    ):
        for info in origen.infolist():
            destino.writestr(info, origen.read(info.filename))
        destino.writestr(nombre, datos)
    return buffer.getvalue()


def test_leer_highly_compressed_zip_raises_too_big_before_opening(
    lector: LectorPlantillaOpenpyxl,
) -> None:
    valida = (PLANTILLAS / PLANTILLA_VALIDA).read_bytes()
    bomba = _con_entrada_extra(valida, "xl/media/relleno.bin", bytes(20 * 1024 * 1024))

    error = motivo_de(lector, bomba)

    assert error.motivo is MotivoPlantillaInvalida.DEMASIADO_GRANDE
    assert str(error) == MENSAJE_DEMASIADO_GRANDE


def test_leer_uncompressed_total_above_limit_raises_too_big(
    lector: LectorPlantillaOpenpyxl, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(lector_plantilla, "MAX_BYTES_DESCOMPRIMIDOS", 100_000)

    error = motivo_de(lector, (PLANTILLAS / PLANTILLA_VALIDA).read_bytes())

    assert error.motivo is MotivoPlantillaInvalida.DEMASIADO_GRANDE


def test_uncompressed_limit_is_fifty_megabytes() -> None:
    assert lector_plantilla.MAX_BYTES_DESCOMPRIMIDOS == 50 * 1024 * 1024


def test_leer_more_iterated_rows_than_limit_raises_too_big(lector: LectorPlantillaOpenpyxl) -> None:
    filas: list[list[object]] = [list(TITULOS_PLANTILLA)]
    filas += [[None] * 14 for _ in range(MAX_FILAS_ITERADAS)]
    filas.append(fila_plantilla(0))

    error = motivo_de(lector, libro_en_bytes({"Registros": filas}))

    assert error.motivo is MotivoPlantillaInvalida.DEMASIADO_GRANDE


def test_leer_rows_up_to_iterated_limit_with_few_data_rows_is_accepted(
    lector: LectorPlantillaOpenpyxl,
) -> None:
    filas: list[list[object]] = [list(TITULOS_PLANTILLA)]
    filas += [[None] * 14 for _ in range(MAX_FILAS_ITERADAS - 1)]
    filas.append(fila_plantilla(0))

    plantilla = lector.leer(libro_en_bytes({"Registros": filas}))

    assert len(plantilla.filas) == 1
    assert plantilla.filas[0].fila == MAX_FILAS_ITERADAS + 1


def test_valid_fixtures_are_far_below_compression_limits() -> None:
    for ruta in PLANTILLAS.glob("*.xlsx"):
        with zipfile.ZipFile(ruta) as archivo:
            infos = archivo.infolist()
        descomprimido = sum(i.file_size for i in infos)
        comprimido = sum(i.compress_size for i in infos)
        assert descomprimido / comprimido < lector_plantilla.MAX_RELACION_COMPRESION / 5


# ------------------------------------- T6-OBS-1 (R-013): avisos de openpyxl sin datos de celdas

SERIAL_FUERA_DE_RANGO = 10**10


def _plantilla_con_fecha_fuera_de_rango() -> bytes:
    """Plantilla con una fecha (columna G) cuyo serial no es una fecha valida."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Registros"
    hoja.append(list(TITULOS_PLANTILLA))
    hoja.append(list(fila_plantilla(0)[:14]))
    celda = hoja.cell(row=2, column=7, value=SERIAL_FUERA_DE_RANGO)
    celda.number_format = "yyyy-mm-dd"
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def test_leer_date_outside_limits_emits_no_openpyxl_warning_with_cell_data(
    lector: LectorPlantillaOpenpyxl,
) -> None:
    contenido = _plantilla_con_fecha_fuera_de_rango()

    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        plantilla = lector.leer(contenido)

    assert [str(a.message) for a in avisos] == []
    assert [f.fila for f in plantilla.filas] == [2]


def test_warning_filter_is_scoped_to_the_read(lector: LectorPlantillaOpenpyxl) -> None:
    contenido = _plantilla_con_fecha_fuera_de_rango()

    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        lector.leer(contenido)
        warnings.warn("aviso posterior a la lectura", UserWarning, stacklevel=1)

    assert [str(a.message) for a in avisos] == ["aviso posterior a la lectura"]
