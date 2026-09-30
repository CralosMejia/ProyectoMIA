"""Generador del Excel de la lista priorizada (CU-12), leyendo el archivo de vuelta.

Datos 100 % sinteticos: ningun nombre, celular ni cedula real.
"""

import io
import os
import tempfile
import time
import zipfile
from dataclasses import replace
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.application.listas.ports import ExportadorLista, ItemLista, ListaDetalle
from app.domain.common.enums import NivelRiesgo, Plan
from app.infrastructure.exportacion.excel_lista import (
    COLUMNAS,
    SIN_FACTORES,
    ExcelListaOpenpyxl,
)

pytestmark = pytest.mark.unit

CREADA_EN = datetime(2026, 9, 1, 12, 30, 15, tzinfo=UTC)
NIVELES = (NivelRiesgo.ALTO, NivelRiesgo.MEDIO, NivelRiesgo.BAJO)


def item(i: int, **cambios: object) -> ItemLista:
    base = ItemLista(
        prediccion_id=100 + i,
        membresia_id=200 + i,
        socio_id=300 + i,
        prioridad=i,
        socio_nombre=f"Socio Sintetico {i:02d}",
        celular=f"09{i:08d}",
        cedula=f"09123456{i:02d}",
        plan=Plan.MENSUAL,
        fecha_vencimiento=date(2026, 9, 1 + i),
        monto_cobrado=Decimal("39.50"),
        score=Decimal("0.9") - Decimal(i) / 100,
        nivel_riesgo=NIVELES[i % 3],
        factores_principales=["Plan: MENSUAL (+0,10)", "Monto actual (USD): 40 (+0,02)"],
        dentro_top_k=i <= 3,
    )
    return replace(base, **cambios)  # type: ignore[arg-type]  # cambios de prueba


def lista(items: tuple[ItemLista, ...] | None = None, top_k: int = 3) -> ListaDetalle:
    return ListaDetalle(
        id=42,
        modelo_id=7,
        modelo_version=1,
        fecha_generacion=date(2026, 8, 31),
        creada_en=CREADA_EN,
        horizonte_dias=30,
        top_k=top_k,
        umbral_alto=Decimal("0.70000000"),
        umbral_medio=Decimal("0.50000000"),
        items=items if items is not None else tuple(item(i) for i in range(1, 6)),
    )


def leer(contenido: bytes):  # type: ignore[no-untyped-def]  # openpyxl sin stubs
    return load_workbook(io.BytesIO(contenido))


def test_exporter_implements_the_port() -> None:
    assert isinstance(ExcelListaOpenpyxl(), ExportadorLista)


def test_generar_has_three_sheets_and_top_k_has_at_most_k_rows() -> None:
    libro = leer(ExcelListaOpenpyxl().generar(lista()))

    assert libro.sheetnames == ["top_3", "cartera_completa", "resumen"]
    assert libro["top_3"].max_row == 1 + 3
    assert libro["cartera_completa"].max_row == 1 + 5


def test_generar_top_k_larger_than_list_keeps_all_rows() -> None:
    libro = leer(ExcelListaOpenpyxl().generar(lista(top_k=25)))

    assert libro.sheetnames[0] == "top_25"
    assert libro["top_25"].max_row == 1 + 5


def test_generar_writes_the_13_columns_with_row_values() -> None:
    libro = leer(ExcelListaOpenpyxl().generar(lista()))
    hoja = libro["cartera_completa"]

    encabezados = [c.value for c in hoja[1]]
    assert encabezados == list(COLUMNAS)
    assert len(encabezados) == 13
    fila = [c.value for c in hoja[2]]
    assert fila[0] == 1
    assert fila[1:5] == ["Socio Sintetico 01", "0900000001", "0912345601", "MENSUAL"]
    assert fila[5].date() == date(2026, 9, 2)
    assert fila[6] == 2  # dias para vencer desde la fecha de generacion
    assert Decimal(str(fila[7])) == Decimal("39.5")
    assert Decimal(str(fila[8])) == Decimal("0.89")
    assert fila[9] == "MEDIO"
    assert fila[10] == "Plan: MENSUAL (+0,10); Monto actual (USD): 40 (+0,02)"
    assert fila[11] is None
    assert fila[12] is None


def test_generar_empty_factors_shows_sin_factores() -> None:
    contenido = ExcelListaOpenpyxl().generar(lista((item(1, factores_principales=[]),)))

    assert leer(contenido)["cartera_completa"]["K2"].value == SIN_FACTORES
    assert SIN_FACTORES == "Sin factores de riesgo destacados"


def test_generar_summary_has_parameters_thresholds_and_counts() -> None:
    hoja = leer(ExcelListaOpenpyxl().generar(lista()))["resumen"]
    valores = {fila[0]: fila[1] for fila in hoja.iter_rows(values_only=True)}

    assert valores["Fecha de generacion"].date() == date(2026, 8, 31)
    assert valores["Horizonte (dias)"] == 30
    assert valores["Top K"] == 3
    assert valores["Version del modelo"] == "v1"
    assert Decimal(str(valores["Umbral alto"])) == Decimal("0.7")
    assert Decimal(str(valores["Umbral medio"])) == Decimal("0.5")
    assert valores["Total en cartera"] == 5
    assert valores["Dentro del top K"] == 3
    # niveles de item(i): i % 3 -> 1 MEDIO, 2 BAJO, 3 ALTO, 4 MEDIO, 5 BAJO
    assert (valores["Nivel ALTO"], valores["Nivel MEDIO"], valores["Nivel BAJO"]) == (1, 2, 2)


@pytest.mark.parametrize("peligroso", ['=HYPERLINK("x")', "+5+5", "-2+3", "@SUM(A1)"])
def test_generar_neutralizes_formulas_with_quote_prefix(peligroso: str) -> None:
    items = (item(1, socio_nombre=peligroso, cedula=peligroso, factores_principales=[peligroso]),)
    contenido = ExcelListaOpenpyxl().generar(lista(items))

    hoja = leer(contenido)["cartera_completa"]
    assert hoja["B2"].value == "'" + peligroso
    assert hoja["B2"].data_type == "s"
    assert hoja["D2"].value == "'" + peligroso
    assert hoja["K2"].value == "'" + peligroso
    with zipfile.ZipFile(io.BytesIO(contenido)) as archivo:
        for nombre in archivo.namelist():
            if nombre.startswith("xl/worksheets/"):
                assert b"<f>" not in archivo.read(nombre)
                assert b"<f " not in archivo.read(nombre)


def test_generar_same_list_produces_identical_bytes_with_properties_from_creada_en(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    exportador = ExcelListaOpenpyxl()
    primero = exportador.generar(lista())
    # Otra hora del sistema: ni el zip ni openpyxl pueden colarla en el resultado.
    monkeypatch.setattr(time, "time", lambda: 1_000_000_000.0)
    segundo = exportador.generar(lista())

    assert primero == segundo
    libro = leer(primero)
    assert libro.properties.created == datetime(2026, 9, 1, 12, 30, 15)  # noqa: DTZ001 - openpyxl devuelve UTC sin zona
    assert libro.properties.modified == datetime(2026, 9, 1, 12, 30, 15)  # noqa: DTZ001 - openpyxl devuelve UTC sin zona
    with zipfile.ZipFile(io.BytesIO(primero)) as archivo:
        assert {i.date_time for i in archivo.infolist()} == {(2026, 9, 1, 12, 30, 14)}


def test_generar_different_creada_en_changes_the_bytes() -> None:
    exportador = ExcelListaOpenpyxl()
    otra = replace(lista(), creada_en=datetime(2026, 9, 2, 8, 0, tzinfo=UTC))

    assert exportador.generar(lista()) != exportador.generar(otra)


def test_generar_in_memory_without_temporary_files(monkeypatch: pytest.MonkeyPatch) -> None:
    import openpyxl.worksheet._writer as escritor

    def prohibido(*_: object, **__: object) -> str:
        raise AssertionError("no se deben crear archivos temporales")

    monkeypatch.setattr(escritor, "create_temporary_file", prohibido)
    directorio = Path(tempfile.gettempdir())
    antes = set(os.listdir(directorio))

    contenido = ExcelListaOpenpyxl().generar(lista())

    assert contenido.startswith(b"PK")
    assert set(os.listdir(directorio)) - antes == set()


def test_generar_empty_list_is_valid() -> None:
    libro = leer(ExcelListaOpenpyxl().generar(lista(())))

    assert libro["top_3"].max_row == 1
    assert libro["cartera_completa"].max_row == 1


def test_generar_has_cedula_column_instead_of_correo_in_both_item_sheets() -> None:
    contenido = ExcelListaOpenpyxl().generar(lista())
    libro = leer(contenido)

    for nombre in ("top_3", "cartera_completa"):
        hoja = libro[nombre]
        encabezados = [c.value for c in hoja[1]]
        assert encabezados[3] == "Cédula"
        assert not any("correo" in str(v).lower() for v in encabezados)
        for fila in hoja.iter_rows(values_only=True):
            assert not any("@" in str(v) for v in fila if v is not None)


def test_generar_cedula_is_text_with_leading_zero_and_none_is_empty() -> None:
    items = (item(1, cedula="0912345678"), item(2, cedula=None))
    contenido = ExcelListaOpenpyxl().generar(lista(items))

    hoja = leer(contenido)["cartera_completa"]
    assert hoja["D2"].value == "0912345678"
    assert hoja["D2"].data_type == "s"
    assert hoja["D3"].value is None
