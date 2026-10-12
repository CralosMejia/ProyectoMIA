"""Lector de libros historicos contra los libros sinteticos de T8 (sin base de datos)."""

import logging
import math
import shutil
import warnings
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.application.historial.ports import LectorHistorico, LibrosLeidos
from app.domain.common.exceptions import DomainError
from app.domain.historico.errores import HojaSinEncabezadoError, LibroSinAnioError
from app.infrastructure.historico.alias_columnas import ALIAS_COLUMNAS, MESES
from app.infrastructure.historico.lector_libros import LectorLibrosExcel
from tests.fixtures import esperado
from tests.fixtures import generar_fixtures as gen

pytestmark = pytest.mark.integration

DIR_HISTORICO = Path(__file__).resolve().parents[2] / "fixtures" / "historico"
RUTAS = [DIR_HISTORICO / esperado.LIBRO_2025, DIR_HISTORICO / esperado.LIBRO_2026]
CAMPOS_PERSONALES = ("nombre", "cedula", "celular", "correo")


@pytest.fixture(scope="module")
def leidos() -> LibrosLeidos:
    return LectorLibrosExcel().leer(RUTAS)


def _textos_personales(leidos: LibrosLeidos) -> set[str]:
    textos: set[str] = set()
    for fila in leidos.filas:
        for campo in CAMPOS_PERSONALES:
            valor = getattr(fila, campo)
            if valor is not None and len(str(valor).strip()) >= 5:
                textos.add(str(valor).strip())
    return textos


def test_lector_satisface_el_puerto() -> None:
    assert isinstance(LectorLibrosExcel(), LectorHistorico)


def test_alias_y_meses_copiados_del_notebook() -> None:
    assert MESES["ENERO"] == 1 and MESES["DICIEMBRE"] == 12 and len(MESES) == 12
    assert ALIAS_COLUMNAS["fecha_inicio"] == ("FECHA INSCRIPCION", "FECHA INICIO")
    assert ALIAS_COLUMNAS["monto_cobrado"] == ("MONTO A COBRAR", "MONTO COBRAR")
    assert "FECHA INSCRIP" not in ALIAS_COLUMNAS["fecha_inicio"]


def test_leer_libros_sinteticos_hojas_incluidas_y_omitidas(leidos: LibrosLeidos) -> None:
    assert [libro.nombre for libro in leidos.libros] == [esperado.LIBRO_2025, esperado.LIBRO_2026]
    for libro in leidos.libros:
        assert libro.anio == esperado.ANIO_POR_LIBRO[libro.nombre]
        assert libro.hojas_incluidas == esperado.HOJAS_INCLUIDAS[libro.nombre]
        assert libro.hojas_omitidas == esperado.HOJAS_OMITIDAS[libro.nombre]


def test_leer_libros_sinteticos_filas_por_hoja_coinciden_con_esperado(
    leidos: LibrosLeidos,
) -> None:
    por_hoja = Counter((fila.anio, fila.hoja) for fila in leidos.filas)
    assert dict(por_hoja) == esperado.FILAS_LEIDAS_POR_HOJA
    assert len(leidos.filas) == esperado.FILAS_LEIDAS


def test_leer_fila_es_la_de_excel_y_empieza_bajo_el_encabezado(leidos: LibrosLeidos) -> None:
    for (anio, hoja), encabezado in esperado.FILA_ENCABEZADO.items():
        filas = [f.fila for f in leidos.filas if (f.anio, f.hoja) == (anio, hoja)]
        # encabezado 0-based -> su fila en Excel es encabezado + 1; los datos, la siguiente.
        assert filas[0] == encabezado + 2
        assert filas == list(range(filas[0], filas[0] + len(filas)))


def test_leer_mes_hoja_es_el_primer_dia_del_mes(leidos: LibrosLeidos) -> None:
    for fila in leidos.filas:
        assert fila.mes_hoja == date(fila.anio, MESES[fila.hoja.strip().upper()], 1)


def test_leer_caso_del_tutor_trae_el_vencimiento(leidos: LibrosLeidos) -> None:
    tutor = esperado.CASO_TUTOR
    (fila,) = [
        f
        for f in leidos.filas
        if (f.anio, f.hoja, f.fila) == (tutor["anio"], tutor["hoja"], tutor["fila"])
    ]
    assert isinstance(fila.fecha_fin, datetime)
    assert fila.fecha_fin.date() == tutor["fecha_vencimiento"]


def test_leer_valores_crudos_sin_nan_ni_tipos_de_pandas(leidos: LibrosLeidos) -> None:
    permitidos = (type(None), str, int, float, datetime, date)
    for fila in leidos.filas:
        for campo in (*CAMPOS_PERSONALES, "plan", "fecha_inicio", "fecha_fin", "forma_pago"):
            valor = getattr(fila, campo)
            assert type(valor) in permitidos or isinstance(valor, str), (campo, type(valor))
            assert not (isinstance(valor, float) and math.isnan(valor))
        assert type(fila.monto_cobrado) in permitidos


def test_leer_cedulas_nunca_son_float(leidos: LibrosLeidos) -> None:
    cedulas = [f.cedula for f in leidos.filas if f.cedula is not None]
    assert cedulas
    assert all(isinstance(c, int | str) for c in cedulas)


def test_leer_hoja_sin_alias_de_inicio_entrega_none(leidos: LibrosLeidos) -> None:
    for anio, hoja in esperado.HOJAS_SIN_FECHA_INICIO:
        filas = [f for f in leidos.filas if (f.anio, f.hoja) == (anio, hoja)]
        assert filas and all(f.fecha_inicio is None for f in filas)


def test_leer_formatos_de_fecha_y_monto_llegan_crudos(leidos: LibrosLeidos) -> None:
    fechas = [f.fecha_inicio for f in leidos.filas] + [f.fecha_fin for f in leidos.filas]
    assert any(isinstance(v, datetime) for v in fechas)
    assert any(isinstance(v, int | float) and not isinstance(v, bool) for v in fechas)
    assert any(isinstance(v, str) and "/" in v for v in fechas)
    assert any(isinstance(f.monto_cobrado, str) for f in leidos.filas)


def test_leer_libro_sin_anio_en_el_nombre_lanza_error(tmp_path: Path) -> None:
    ruta = tmp_path / "Control Gym (1).xlsx"
    shutil.copy(RUTAS[0], ruta)
    with pytest.raises(LibroSinAnioError) as error:
        LectorLibrosExcel().leer([ruta])
    assert isinstance(error.value, DomainError)
    assert "Control Gym (1).xlsx" in str(error.value)


def test_leer_hoja_mensual_sin_encabezado_lanza_error_con_libro_y_hoja(tmp_path: Path) -> None:
    ruta = tmp_path / "Control Gym 2024 prueba.xlsx"
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "MAYO "
    hoja.append(["Nombre", "Telefono"])
    hoja.append(["Valentina Robles", 991234567])
    libro.save(ruta)

    with pytest.raises(HojaSinEncabezadoError) as error:
        LectorLibrosExcel().leer([ruta])
    mensaje = str(error.value)
    assert isinstance(error.value, DomainError)
    assert "Control Gym 2024 prueba.xlsx" in mensaje and "MAYO" in mensaje
    assert "Valentina" not in mensaje


def test_leer_encabezado_despues_de_diez_filas_no_se_detecta(tmp_path: Path) -> None:
    ruta = tmp_path / "Control Gym 2024 tarde.xlsx"
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "JUNIO"
    for i in range(10):
        hoja.append([f"nota {i}"])
    hoja.append(["NOMBRES Y APELLIDOS", "CEDULA"])
    libro.save(ruta)

    with pytest.raises(HojaSinEncabezadoError):
        LectorLibrosExcel().leer([ruta])


def test_leer_columna_ausente_sale_none(tmp_path: Path) -> None:
    ruta = tmp_path / "Control Gym 2024 minimo.xlsx"
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "Julio"
    hoja.append(["Nombres y apellidos del socio", "Cédula"])
    hoja.append(["Valentina Robles", "1712345678"])
    libro.save(ruta)

    (fila,) = LectorLibrosExcel().leer([ruta]).filas
    assert (fila.anio, fila.hoja, fila.fila, fila.mes_hoja) == (2024, "Julio", 2, date(2024, 7, 1))
    assert fila.cedula == "1712345678"
    assert fila.plan is None and fila.fecha_fin is None and fila.monto_cobrado is None


def test_leer_no_escribe_logs_con_datos_de_celdas(
    leidos: LibrosLeidos, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level(logging.DEBUG):
        LectorLibrosExcel().leer(RUTAS)
    personales = _textos_personales(leidos) | set(gen.NOMBRES_FICTICIOS)
    assert personales
    for registro in caplog.records:
        mensaje = registro.getMessage()
        assert not any(texto in mensaje for texto in personales)


def test_leer_no_escribe_en_disco(tmp_path: Path) -> None:
    for ruta in RUTAS:
        shutil.copy(ruta, tmp_path / ruta.name)
    antes = {p: p.stat().st_mtime_ns for p in tmp_path.iterdir()}
    LectorLibrosExcel().leer(sorted(tmp_path.iterdir()))
    assert {p: p.stat().st_mtime_ns for p in tmp_path.iterdir()} == antes


# ------------------------------------------------------------------ casos limite (test-engineer)


def _libro(ruta: Path, hojas: dict[str, list[list[object]]]) -> Path:
    libro = Workbook()
    primera = libro.active
    assert primera is not None
    libro.remove(primera)
    for titulo, filas in hojas.items():
        hoja = libro.create_sheet(titulo)
        for fila in filas:
            hoja.append(fila)
    libro.save(ruta)
    return ruta


@pytest.mark.parametrize(
    ("nombre", "anio"),
    [
        ("Control Gym 2025.xlsx", 2025),
        ("Control Gym 2026 sintetico.xlsx", 2026),
        ("Control_Gym_2024_v2.xlsx", 2024),
    ],
)
def test_leer_anio_del_nombre_del_libro_se_toma_del_nombre(
    tmp_path: Path, nombre: str, anio: int
) -> None:
    ruta = _libro(tmp_path / nombre, {"ENERO": [["NOMBRE", "CEDULA"]]})
    (libro,) = LectorLibrosExcel().leer([ruta]).libros
    assert libro.anio == anio


@pytest.mark.parametrize("nombre", ["Control Gym (1).xlsx", "Control Gym 20255.xlsx", "Gym 1.xlsx"])
def test_leer_nombre_sin_anio_de_cuatro_digitos_lanza_libro_sin_anio(
    tmp_path: Path, nombre: str
) -> None:
    ruta = _libro(tmp_path / nombre, {"ENERO": [["NOMBRE", "CEDULA"]]})
    with pytest.raises(LibroSinAnioError):
        LectorLibrosExcel().leer([ruta])


def test_leer_hojas_mes_en_minusculas_y_con_espacios_se_incluyen_y_otras_se_omiten(
    tmp_path: Path,
) -> None:
    encabezado: list[list[object]] = [["NOMBRE", "CEDULA"], ["Valentina Robles", "1712345678"]]
    ruta = _libro(
        tmp_path / "Control Gym 2024 hojas.xlsx",
        {" agosto ": encabezado, "AGOSTO 2": encabezado, "Setiembre": encabezado},
    )
    leidos = LectorLibrosExcel().leer([ruta])
    (libro,) = leidos.libros
    assert libro.hojas_incluidas == (" agosto ",)
    assert libro.hojas_omitidas == ("AGOSTO 2", "Setiembre")
    assert [f.mes_hoja for f in leidos.filas] == [date(2024, 8, 1)]


def test_leer_encabezado_en_la_decima_fila_se_detecta(tmp_path: Path) -> None:
    filas: list[list[object]] = [[f"nota {i}"] for i in range(9)]
    filas += [["NOMBRES Y APELLIDOS", "CEDULA"], ["Valentina Robles", "1712345678"]]
    ruta = _libro(tmp_path / "Control Gym 2024 limite.xlsx", {"JUNIO": filas})
    (fila,) = LectorLibrosExcel().leer([ruta]).filas
    assert fila.fila == 11
    assert fila.cedula == "1712345678"


def test_leer_encabezado_con_solo_cedula_o_solo_nombre_lanza_error(tmp_path: Path) -> None:
    ruta = _libro(
        tmp_path / "Control Gym 2024 parcial.xlsx",
        {"MAYO": [["CEDULA", "PLAN"], ["NOMBRE", "PLAN"], ["x", "y"]]},
    )
    with pytest.raises(HojaSinEncabezadoError):
        LectorLibrosExcel().leer([ruta])


def test_leer_valores_crudos_conservan_su_tipo_exacto(tmp_path: Path) -> None:
    ruta = _libro(
        tmp_path / "Control Gym 2024 tipos.xlsx",
        {
            "ENERO": [
                ["NOMBRE", "CEDULA", "PLAN", "FECHA INICIO", "FECHA FIN", "MONTO A COBRAR"],
                ["Valentina Robles", "0912345678", "1 MES", datetime(2024, 1, 5), 45352, 39.5],  # noqa: DTZ001 - celda Excel sin zona
                [None, None, None, None, None, None],  # vacia intermedia
                ["Tomas Aguirre", 1712345678, None, "05/01/2024", None, 25],
            ]
        },
    )
    primera, vacia, segunda = LectorLibrosExcel().leer([ruta]).filas
    assert primera.cedula == "0912345678"  # el cero inicial no se pierde
    assert type(primera.fecha_inicio) is datetime
    assert primera.fecha_inicio == datetime(2024, 1, 5)  # noqa: DTZ001 - celda Excel sin zona
    assert type(primera.fecha_fin) is int and primera.fecha_fin == 45352
    assert type(primera.monto_cobrado) is float and primera.monto_cobrado == 39.5
    assert type(segunda.cedula) is int and segunda.cedula == 1712345678
    assert segunda.fecha_inicio == "05/01/2024"
    assert type(segunda.monto_cobrado) is int and segunda.monto_cobrado == 25
    assert segunda.plan is None and segunda.fecha_fin is None
    assert all(
        getattr(vacia, campo) is None
        for campo in (*CAMPOS_PERSONALES, "plan", "fecha_inicio", "fecha_fin", "monto_cobrado")
    )


def test_leer_error_de_una_hoja_no_devuelve_resultado_parcial(tmp_path: Path) -> None:
    buena = _libro(
        tmp_path / "Control Gym 2023 bueno.xlsx",
        {"ENERO": [["NOMBRE", "CEDULA"], ["Valentina Robles", "1712345678"]]},
    )
    mala = _libro(tmp_path / "Control Gym 2024 malo.xlsx", {"MAYO": [["sin", "titulos"]]})
    with pytest.raises(HojaSinEncabezadoError) as error:
        LectorLibrosExcel().leer([buena, mala])
    assert (error.value.libro, error.value.hoja) == ("Control Gym 2024 malo.xlsx", "MAYO")


def test_leer_fecha_fuera_de_rango_no_emite_avisos_de_openpyxl_con_datos_de_celdas(
    tmp_path: Path,
) -> None:
    ruta = _libro(
        tmp_path / "Control Gym 2024.xlsx",
        {"MAYO": [["CEDULA", "NOMBRE", "FECHA INICIO"], ["1700000019", "Socio Ficticio", None]]},
    )
    libro = load_workbook(ruta)
    celda = libro["MAYO"]["C2"]
    celda.value = 10**10
    celda.number_format = "yyyy-mm-dd"
    libro.save(ruta)

    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        leidos = LectorLibrosExcel().leer([ruta])

    assert [str(a.message) for a in avisos] == []
    assert [f.fila for f in leidos.filas] == [2]
