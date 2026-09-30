"""Descarga del Excel de una lista (CU-12) contra PostgreSQL real.

Aislamiento entre usuarios: NO APLICA. No hay login (RT-12): el acceso se controla por IP
(CU-00) y no existe un propietario de datos por usuario.
"""

import io
from collections.abc import Sequence

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine

from app.api.dependencies import get_predictor
from app.domain.common.predictor import FeatureRow, Prediction, TrainedModel
from tests.integration.api.listas_datos import (
    FILAS_ESPERADAS,
    crear_modelo_vigente,
    fijar_reloj_api,
)

pytestmark = pytest.mark.integration

LISTAS = "/api/v1/listas"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _reloj_fijo(contract_app: FastAPI) -> None:
    fijar_reloj_api(contract_app)


class PredictorEspia:
    def __init__(self) -> None:
        self.llamadas = 0

    def train(self, *args: object, **kwargs: object) -> TrainedModel:
        raise AssertionError("no se entrena")

    def predict(self, artefacto: bytes, features: Sequence[FeatureRow]) -> list[Prediction]:
        self.llamadas += 1
        raise AssertionError("la descarga no debe ejecutar el modelo")


def test_get_excel_returns_xlsx_with_safe_filename(
    contract_client: TestClient, base_migrada: Engine
) -> None:
    lista = contract_client.post(
        LISTAS, json={"fecha_referencia": "2026-08-31", "horizonte_dias": 30, "top_k": 25}
    ).json()

    respuesta = contract_client.get(f"{LISTAS}/{lista['id']}/excel")

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == XLSX
    assert (
        respuesta.headers["content-disposition"]
        == f'attachment; filename="lista_{lista["id"]}_2026-08-31.xlsx"'
    )
    libro = load_workbook(io.BytesIO(respuesta.content))
    assert libro.sheetnames == ["top_25", "cartera_completa", "resumen"]
    assert libro["top_25"].max_row == 1 + 25
    assert libro["cartera_completa"].max_row == 1 + FILAS_ESPERADAS


def test_get_excel_twice_with_new_model_is_identical_and_does_not_predict(
    contract_app: FastAPI, contract_client: TestClient, base_migrada: Engine
) -> None:
    lista = contract_client.post(
        LISTAS, json={"fecha_referencia": "2026-08-31", "horizonte_dias": 30, "top_k": 25}
    ).json()
    espia = PredictorEspia()
    contract_app.dependency_overrides[get_predictor] = lambda: espia
    try:
        primera = contract_client.get(f"{LISTAS}/{lista['id']}/excel")
        crear_modelo_vigente(base_migrada, version=2)
        segunda = contract_client.get(f"{LISTAS}/{lista['id']}/excel")
    finally:
        contract_app.dependency_overrides.pop(get_predictor, None)

    assert primera.status_code == segunda.status_code == 200
    assert primera.content == segunda.content
    assert espia.llamadas == 0
    resumen = load_workbook(io.BytesIO(segunda.content))["resumen"]
    valores = {fila[0]: fila[1] for fila in resumen.iter_rows(values_only=True)}
    assert valores["Version del modelo"] == "v1"


@pytest.mark.parametrize(("ruta", "codigo"), [("/999999/excel", 404), ("/abc/excel", 400)])
def test_get_excel_unknown_or_invalid_id(
    contract_client: TestClient, base_limpia: Engine, ruta: str, codigo: int
) -> None:
    respuesta = contract_client.get(f"{LISTAS}{ruta}")

    assert respuesta.status_code == codigo
    assert set(respuesta.json()) == {"error"}


def test_get_excel_has_cedula_column_as_text_and_no_correo(
    contract_client: TestClient, base_migrada: Engine
) -> None:
    lista = contract_client.post(
        LISTAS, json={"fecha_referencia": "2026-08-31", "horizonte_dias": 30, "top_k": 25}
    ).json()

    excel = contract_client.get(f"{LISTAS}/{lista['id']}/excel")
    libro = load_workbook(io.BytesIO(excel.content))

    for nombre in ("top_25", "cartera_completa"):
        hoja = libro[nombre]
        assert hoja["D1"].value == "Cédula"
        valores = [v for fila in hoja.iter_rows(values_only=True) for v in fila if v is not None]
        assert not any("@" in str(v) or "correo" in str(v).lower() for v in valores)
        assert all(hoja.cell(row=r, column=4).data_type == "s" for r in range(2, hoja.max_row + 1))
