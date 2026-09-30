"""API de cargas (CU-01): fronteras, concurrencia y datos personales (R-007 T10, test-engineer).

Aislamiento entre usuarios: NO APLICA. No hay login (RT-12): el acceso se controla por IP
(CU-00) y no existe un propietario de datos por usuario.

Complementa ``test_cargas_api.py`` con los criterios que alli no se cubren:
- archivo ilegible con la migracion pendiente -> ``400`` (se valida el archivo antes que el estado);
- frontera de tamano: exactamente 5 MB no es ``413``;
- ``400`` por hoja ``Registros`` ausente y por titulos distintos (mensajes del contrato);
- cedula con cero inicial enlaza con la historica guardada sin el cero (decision (d));
- dos confirmaciones simultaneas del mismo archivo: una ``201`` y otra ``422``, sin duplicar;
- ni los logs ni las respuestas contienen nombres, cedulas, celulares ni correos.

Datos 100 % sinteticos (fixtures y libros en ``tmp_path`` con valores ficticios).
"""

import io
import logging
import threading
import time
from collections.abc import Iterator
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook  # type: ignore[import-untyped]  # openpyxl no publica stubs
from sqlalchemy import Engine, text
from sqlalchemy.orm import Session

from app.application.historial.ports import NuevoSocio
from app.domain.historico.normalizacion import normalizar_nombre
from app.infrastructure.cargas.lector_plantilla import TITULOS_PLANTILLA
from app.infrastructure.persistence.repositories.historial import SqlAlchemySocioRepository
from app.infrastructure.persistence.unit_of_work_cargas import CLAVE_BLOQUEO_CARGAS
from tests.fixtures.generar_fixtures import (
    DIR_PLANTILLAS,
    PLANTILLA_CEDULA_INVALIDA,
    PLANTILLA_VALIDA,
    fila_plantilla,
    filas_plantillas,
    generar_plantilla,
)
from tests.integration.api.listas_datos import RELOJ_API, fijar_reloj_api

pytestmark = pytest.mark.integration

VISTA_PREVIA = "/api/v1/cargas/vista-previa"
CARGAS = "/api/v1/cargas"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_BYTES = 5 * 1024 * 1024
MENSAJE_DANADO = "El archivo está dañado o no es un .xlsx válido."
_CONTEOS = (
    "SELECT (SELECT count(*) FROM socios), (SELECT count(*) FROM membresias), "
    "(SELECT count(*) FROM cargas)"
)


class Reloj:
    def now(self) -> datetime:
        return RELOJ_API


@pytest.fixture
def cliente(contract_app: FastAPI, base_migrada: Engine) -> Iterator[TestClient]:
    fijar_reloj_api(contract_app)
    with TestClient(contract_app, raise_server_exceptions=False) as test_client:
        yield test_client


def _archivo(contenido: bytes, nombre: str = "plantilla.xlsx") -> dict[str, Any]:
    return {"archivo": (nombre, contenido, XLSX)}


def _fixture(nombre: str) -> bytes:
    return (DIR_PLANTILLAS / nombre).read_bytes()


def _conteos(engine: Engine) -> dict[str, int]:
    with engine.connect() as conexion:
        fila = conexion.execute(text(_CONTEOS)).one()
    return {"socios": int(fila[0]), "membresias": int(fila[1]), "cargas": int(fila[2])}


def _libro(hojas: dict[str, list[list[object]]]) -> bytes:
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for numero, valores in enumerate(filas, start=1):
            for columna, valor in enumerate(valores, start=1):
                if valor is not None:
                    hoja.cell(row=numero, column=columna, value=valor)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _cedula_con_cero(secuencia: int) -> str:
    """Cedula ficticia ``09XXXXXXXV`` con verificador modulo 10 valido."""
    base = f"09{secuencia:07d}"
    total = 0
    for posicion, caracter in enumerate(base):
        producto = int(caracter) * (2 if posicion % 2 == 0 else 1)
        total += producto - 9 if producto > 9 else producto
    return f"{base}{(10 - total % 10) % 10}"


# ---------------------------------------------------------------- orden de validacion


@pytest.mark.parametrize("path", [VISTA_PREVIA, CARGAS])
def test_post_unreadable_file_with_migration_pending_returns_400_not_409(
    contract_app: FastAPI, base_limpia: Engine, path: str
) -> None:
    fijar_reloj_api(contract_app)
    with TestClient(contract_app, raise_server_exceptions=False) as cliente:
        response = cliente.post(path, files=_archivo(b"texto plano renombrado", "notas.xlsx"))

    assert response.status_code == 400
    assert response.json() == {"error": {"code": "validation_error", "message": MENSAJE_DANADO}}
    assert _conteos(base_limpia) == {"socios": 0, "membresias": 0, "cargas": 0}


# ---------------------------------------------------------------- frontera de tamano


@pytest.mark.parametrize("path", [VISTA_PREVIA, CARGAS])
def test_post_file_of_exactly_5_mb_is_not_rejected_as_too_large(
    cliente: TestClient, path: str
) -> None:
    response = cliente.post(path, files=_archivo(b"0" * MAX_BYTES))

    assert response.status_code == 400
    assert response.json() == {"error": {"code": "validation_error", "message": MENSAJE_DANADO}}


# ---------------------------------------------------------------- estructura del libro


@pytest.mark.parametrize("path", [VISTA_PREVIA, CARGAS])
@pytest.mark.parametrize(
    ("hojas", "mensaje"),
    [
        (
            {"Hoja1": [list(TITULOS_PLANTILLA), fila_plantilla(0)]},
            "El archivo no tiene la hoja Registros.",
        ),
        (
            {"Registros": [[*TITULOS_PLANTILLA[:5], "OTRO", *TITULOS_PLANTILLA[6:]]]},
            "Los títulos de la hoja Registros no coinciden con la plantilla oficial.",
        ),
    ],
    ids=["sin-registros", "titulos-distintos"],
)
def test_post_wrong_structure_returns_400_with_contract_message(
    cliente: TestClient,
    base_migrada: Engine,
    path: str,
    hojas: dict[str, list[list[object]]],
    mensaje: str,
) -> None:
    antes = _conteos(base_migrada)

    response = cliente.post(path, files=_archivo(_libro(hojas)))

    assert response.status_code == 400
    assert response.json() == {"error": {"code": "validation_error", "message": mensaje}}
    assert _conteos(base_migrada) == antes


# ---------------------------------------------------------------- cedula con cero inicial


def test_confirm_cedula_with_leading_zero_links_historic_member_saved_without_zero(
    cliente: TestClient, base_migrada: Engine, tmp_path: Path
) -> None:
    cedula = _cedula_con_cero(4321)
    nombre_historico = "Ficticio Historico Cero"
    with Session(base_migrada) as session, session.begin():
        [socio_id] = SqlAlchemySocioRepository(session, Reloj()).insertar_en_bloque(
            [NuevoSocio(cedula[1:], nombre_historico, normalizar_nombre(nombre_historico))]
        )
    fila = fila_plantilla(0)
    fila[1] = "Ficticio Nombre Distinto"
    fila[2] = cedula
    antes = _conteos(base_migrada)

    response = cliente.post(
        CARGAS, files=_archivo(generar_plantilla([fila], tmp_path / "p.xlsx").read_bytes())
    )

    assert response.status_code == 201
    assert response.json()["socios_enlazados"] == 1
    assert response.json()["filas_guardadas"] == 1
    assert _conteos(base_migrada)["socios"] == antes["socios"]
    with base_migrada.connect() as conexion:
        duenos = conexion.execute(
            text("SELECT socio_id FROM membresias WHERE carga_id = :id"),
            {"id": response.json()["id"]},
        ).scalars()
        assert list(duenos) == [socio_id]


# ---------------------------------------------------------------- concurrencia


def _esperar_bloqueados(engine: Engine, cuantos: int, limite_s: float = 15.0) -> int:
    """Espera a que ``cuantos`` sesiones esperen el bloqueo consultivo de cargas."""
    fin = time.monotonic() + limite_s
    esperando = 0
    while time.monotonic() < fin:
        with engine.connect() as conexion:
            esperando = int(
                conexion.execute(
                    text(
                        "SELECT count(*) FROM pg_locks WHERE locktype = 'advisory' AND NOT granted"
                    )
                ).scalar_one()
            )
        if esperando >= cuantos:
            return esperando
        time.sleep(0.05)
    return esperando


def test_confirm_same_file_twice_concurrently_saves_once_and_rejects_the_other(
    cliente: TestClient, base_migrada: Engine
) -> None:
    contenido = _fixture(PLANTILLA_VALIDA)
    antes = _conteos(base_migrada)
    respuestas: list[httpx.Response] = []
    candado = threading.Lock()

    def confirmar() -> None:
        response = cliente.post(CARGAS, files=_archivo(contenido))
        with candado:
            respuestas.append(response)

    hilos = [threading.Thread(target=confirmar, daemon=True) for _ in range(2)]
    # El test retiene el bloqueo para que ambas confirmaciones esten en vuelo a la vez.
    with base_migrada.connect() as retenedor, retenedor.begin():
        retenedor.execute(
            text("SELECT pg_advisory_xact_lock(:clave)"), {"clave": CLAVE_BLOQUEO_CARGAS}
        )
        for hilo in hilos:
            hilo.start()
        assert _esperar_bloqueados(base_migrada, 2) == 2
    for hilo in hilos:
        hilo.join(timeout=30)

    assert sorted(r.status_code for r in respuestas) == [201, 422]
    [rechazada] = [r for r in respuestas if r.status_code == 422]
    assert rechazada.json() == {
        "error": {
            "code": "domain_rule_violated",
            "message": "El archivo no tiene filas nuevas para guardar.",
        }
    }
    despues = _conteos(base_migrada)
    assert despues["membresias"] == antes["membresias"] + 6
    assert despues["cargas"] == antes["cargas"] + 1


# ---------------------------------------------------------------- datos personales


def _datos_personales(archivo: str) -> list[str]:
    valores: list[str] = []
    for fila in filas_plantillas()[archivo]:
        valores.extend(str(v) for v in fila[1:5] if v not in (None, ""))
    return valores


def test_preview_and_confirm_never_expose_personal_data_in_logs_or_responses(
    cliente: TestClient, caplog: pytest.LogCaptureFixture
) -> None:
    personales = _datos_personales(PLANTILLA_VALIDA) + _datos_personales(PLANTILLA_CEDULA_INVALIDA)

    assert personales
    with caplog.at_level(logging.DEBUG):
        respuestas = [
            cliente.post(VISTA_PREVIA, files=_archivo(_fixture(PLANTILLA_CEDULA_INVALIDA))),
            cliente.post(CARGAS, files=_archivo(_fixture(PLANTILLA_CEDULA_INVALIDA))),
            cliente.post(VISTA_PREVIA, files=_archivo(_fixture(PLANTILLA_VALIDA))),
            cliente.post(CARGAS, files=_archivo(_fixture(PLANTILLA_VALIDA))),
            cliente.post(CARGAS, files=_archivo(_fixture(PLANTILLA_VALIDA))),
        ]

    assert [r.status_code for r in respuestas] == [200, 422, 200, 201, 422]
    textos = caplog.text + "".join(r.text for r in respuestas)
    filtrados = [valor for valor in personales if valor in textos]
    assert filtrados == []
